#!/usr/bin/env python3
"""Render the user guide to DOCX and self-contained HTML without resizing assets.

Run with Python containing python-docx and Pillow. PDF is exported from the DOCX
by build_manual.sh, so the two printable formats share the same pagination.
"""
from __future__ import annotations

import base64
import html
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/使用说明.md'
FONT = 'Arial Unicode MS'
TOKEN = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def inline_html(text):
    parts = TOKEN.split(text)
    out = []
    for part in parts:
        if part.startswith('**'):
            out.append('<strong>' + html.escape(part[2:-2]) + '</strong>')
        elif part.startswith('`'):
            out.append('<code>' + html.escape(part[1:-1]) + '</code>')
        elif part.startswith('['):
            m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
            out.append('<a href="%s">%s</a>' % (html.escape(m[2], quote=True), html.escape(m[1])))
        else:
            out.append(html.escape(part))
    return ''.join(out)


def font(run, size=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    if size:
        run.font.size = Pt(size)


def inline_word(paragraph, text):
    for part in TOKEN.split(text):
        if part.startswith('['):
            m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(
                m[2], 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
            run = OxmlElement('w:r')
            pr = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '1E78FA'); pr.append(color)
            fonts = OxmlElement('w:rFonts')
            for attr in ('ascii', 'hAnsi', 'eastAsia'):
                fonts.set(qn('w:' + attr), FONT)
            pr.append(fonts); run.append(pr)
            value = OxmlElement('w:t'); value.text = m[1]; run.append(value)
            link.append(run); paragraph._p.append(link)
        else:
            bold = part.startswith('**')
            code = part.startswith('`')
            run = paragraph.add_run(part[2:-2] if bold else part[1:-1] if code else part)
            run.bold = bold
            font(run, 10 if code else None)


def parse():
    lines = SOURCE.read_text().splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line:
            continue
        if line == '<!-- pagebreak -->':
            blocks.append(('break', None)); continue
        if line.startswith('```'):
            code = []
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i += 1
            i += 1; blocks.append(('code', '\n'.join(code))); continue
        if line.startswith('|'):
            rows = [[x.strip() for x in line.strip('|').split('|')]]
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', x) for x in row):
                    rows.append(row)
                i += 1
            blocks.append(('table', rows)); continue
        image = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if image:
            blocks.append(('image', (image[1], image[2]))); continue
        heading = re.match(r'^(#{1,3}) (.*)', line)
        if heading:
            blocks.append(('heading', (len(heading[1]), heading[2]))); continue
        blocks.append(('paragraph', line))
    return blocks


def build():
    doc = Document()
    doc.core_properties.title = 'SlideCraft 使用手册'
    doc.core_properties.author = 'SlideCraft'
    doc.core_properties.subject = '安装部署与操作指南'
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(.55); sec.bottom_margin = Inches(.55)
    sec.left_margin = Inches(.7); sec.right_margin = Inches(.7)
    for name, size in [('Normal', 10.5), ('Title', 24), ('Heading 1', 18), ('Heading 2', 13), ('Caption', 9)]:
        style = doc.styles[name]
        style.font.name = FONT; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.space_before = Pt(0 if name in ('Normal', 'Title', 'Caption') else 7)
        style.paragraph_format.keep_with_next = name in ('Title', 'Heading 1', 'Heading 2')
        for border in list(style._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        style.paragraph_format.line_spacing = 1.0
    doc.styles['Normal'].paragraph_format.widow_control = True
    settings = doc.settings.element
    no_compress = OxmlElement('w:doNotAutoCompressPictures'); settings.append(no_compress)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    fragments = []
    page = 1
    new_page = False
    for kind, data in parse():
        if kind == 'break':
            new_page = True; fragments.append('<div class="pagebreak"></div>'); page += 1
        elif kind == 'heading':
            level, title = data
            p = doc.add_paragraph(title, 'Title' if level == 1 else 'Heading 1' if level == 2 else 'Heading 2')
            if new_page:
                p.paragraph_format.page_break_before = True
                new_page = False
            fragments.append(f'<h{level}>{html.escape(title)}</h{level}>')
        elif kind == 'paragraph':
            p = doc.add_paragraph(); inline_word(p, data)
            fragments.append('<p>' + inline_html(data) + '</p>')
        elif kind == 'code':
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.keep_together = True
            r = p.add_run(data); font(r, 9)
            fragments.append('<pre><code>' + html.escape(data) + '</code></pre>')
        elif kind == 'image':
            alt, relative = data
            path = SOURCE.parent / relative
            with Image.open(path) as im:
                w, h = im.size
            max_height = 1.65 if page in (6, 10) else 3.5 if page in (7, 8, 9) else 2.05 if page == 4 else 2.35 if 'screenshots/' in relative else 2.55
            width = min(7.1, max_height * w / h)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            r = p.add_run(); r.add_picture(str(path), width=Inches(width))
            for node in r._element.iter(qn('wp:docPr')):
                node.set('descr', alt)
            cap = doc.add_paragraph(alt, 'Caption'); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            encoded = base64.b64encode(path.read_bytes()).decode()
            fragments.append('<figure><img src="data:image/png;base64,%s" alt="%s"><figcaption>%s</figcaption></figure>' %
                             (encoded, html.escape(alt, quote=True), html.escape(alt)))
        elif kind == 'table':
            rows = data
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.autofit = False
            widths = [2.1, 2, 3] if len(rows[0]) == 3 else [3.1, 4]
            for j, width in enumerate(widths):
                table.columns[j].width = Inches(width)
            for n, row in enumerate(rows):
                cells = table.add_row().cells
                for j, value in enumerate(row):
                    cell = cells[j]; cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]; inline_word(p, value); p.paragraph_format.space_after = Pt(5); p.paragraph_format.space_before = Pt(5)
                    tcpr = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement('w:tcMar')
                    for side in ('top', 'bottom', 'left', 'right'):
                        m = OxmlElement('w:' + side); m.set(qn('w:w'), '80'); m.set(qn('w:type'), 'dxa'); margins.append(m)
                    tcpr.append(margins)
                    borders = OxmlElement('w:tcBorders')
                    for side in ('top', 'bottom', 'left', 'right'):
                        border = OxmlElement('w:' + side)
                        for attr, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                            border.set(qn('w:' + attr), val)
                        borders.append(border)
                    tcpr.append(borders)
                    if n == 0:
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAF2FF'); tcpr.append(shade)
                        for r in p.runs:
                            r.bold = True
                if n == 0:
                    repeat = OxmlElement('w:tblHeader'); table.rows[-1]._tr.get_or_add_trPr().append(repeat)
            fragments.append('<table>' + ''.join('<tr>' + ''.join(
                f'<{"th" if n == 0 else "td"}>{inline_html(x)}</{"th" if n == 0 else "td"}>' for x in row) + '</tr>'
                for n, row in enumerate(rows)) + '</table>')
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
    doc.save(SOURCE.with_suffix('.docx'))
    css = '''body{max-width:920px;margin:40px auto;padding:0 24px;color:#182238;font:16px/1.7 "PingFang SC","Microsoft YaHei",sans-serif;background:#fff}
h1,h2,h3{color:#000;line-height:1.3}h2{margin-top:38px}a{color:#1e78fa}figure{margin:24px 0;text-align:center}img{max-width:100%;height:auto}figcaption{font-size:13px;color:#536077}
table{width:100%;border-collapse:collapse;margin:16px 0}th,td{border:1px solid #d9d9d9;padding:10px;text-align:left}th{background:#eaf2ff}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:14px;padding:0}code{font-family:ui-monospace,"PingFang SC",monospace}
@media print{@page{size:letter;margin:14mm 18mm}body{font-size:11pt;margin:0;padding:0}.pagebreak{break-before:page}figure,table,pre{break-inside:avoid}h2,h3{break-after:avoid}}'''
    SOURCE.with_suffix('.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>SlideCraft 使用手册</title><style>' + css + '</style><body>' + '\n'.join(fragments) + '</body></html>')
    # Keep Markdown and its referenced PNGs together for offline reading.
    assets = [data[1] for kind, data in parse() if kind == 'image']
    with zipfile.ZipFile(SOURCE.parent / 'SlideCraft-使用手册-Markdown.zip', 'w', zipfile.ZIP_DEFLATED) as bundle:
        bundle.write(SOURCE, SOURCE.name)
        for relative in sorted(set(assets)):
            bundle.write(SOURCE.parent / relative, relative)
    print(f'Built {page} sections; original-resolution images retained in DOCX, HTML and Markdown bundle.')


if __name__ == '__main__':
    build()
