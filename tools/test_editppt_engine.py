"""高精度引擎（第三方 image-to-editable-ppt 接入）离线回归：不跑 CLI、不联网、不花额度。

锁定的是接入层最容易出错的地方：
- 可用性判断（CLI/skill/agent 三件套缺一不可，缺了要说清缺什么）
- 提示合成（文字必须来自识别结果不改写；几何优先用墨迹测量；距离过远不硬套 size_group）
- 校验结果判定（缺文件/坏 JSON/未通过都不能当成成功）
- 失败聚合（任何一页没过校验 → 整体 ok=False 且保留现场，绝不伪造成功）
"""
import json
import os
import shutil
import sys
import tempfile
import types

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)

from server.app import editppt_engine as ee  # noqa: E402

PASS, FAIL = [], []


def check(name, ok, detail=""):
    (PASS if ok else FAIL).append(name)
    print("  %s %s%s" % ("✅" if ok else "❌", name, (" — " + detail) if detail else ""))


class Fake:
    """保存并恢复被 monkeypatch 的模块属性。"""

    def __init__(self):
        self.saved = []

    def set(self, obj, name, value):
        self.saved.append((obj, name, getattr(obj, name)))
        setattr(obj, name, value)

    def restore(self):
        for obj, name, value in reversed(self.saved):
            setattr(obj, name, value)
        self.saved = []


def main():
    tmp = tempfile.mkdtemp(prefix="editppt-test-")
    fk = Fake()
    try:
        print("[1] 可用性：三件套缺一不可，且要说清缺什么")
        fk.set(ee, "cli_path", lambda: os.path.join(tmp, "nope", "editppt"))
        st = ee.available()
        check("CLI 缺失 → ok=False 且提示未安装", st["ok"] is False and "未安装" in st["reason"], st["reason"])
        cli = os.path.join(tmp, "bin", "editppt")       # 放在 bin/ 下，避免与运行根目录（FILES_DIR/editppt）撞名
        os.makedirs(os.path.dirname(cli), exist_ok=True)
        open(cli, "w").write("#!/bin/sh\n")
        os.chmod(cli, 0o755)
        fk.set(ee, "cli_path", lambda: cli)
        fk.set(ee, "skill_dir", lambda: os.path.join(tmp, "noskill"))
        st = ee.available()
        check("skill 目录缺失 → 指出来", st["ok"] is False and "skill" in st["reason"], st["reason"])
        skill = os.path.join(tmp, "skill")
        os.makedirs(skill)
        fk.set(ee, "skill_dir", lambda: skill)
        st = ee.available()
        check("skill 目录无 SKILL.md → 指出来", st["ok"] is False and "SKILL.md" in st["reason"], st["reason"])
        open(os.path.join(skill, "SKILL.md"), "w").write("# x")
        fk.set(ee, "agent_kind", lambda: "hermes")     # 本用例验 Hermes 兜底执行器的缺件提示
        fk.set(ee, "hermes_path", lambda: os.path.join(tmp, "no-hermes"))
        # 本机若装了 hermes，available() 会经 shutil.which 在 PATH 里兜底找到它；
        # 该断言要验的是「三件套缺 agent 时说清缺什么」，故只在 ee 模块内隔离 PATH 查找，不动全局。
        _real_shutil = ee.shutil
        ee.shutil = types.SimpleNamespace(which=lambda *a, **k: None)
        try:
            st = ee.available()
        finally:
            ee.shutil = _real_shutil
        check("agent 缺失 → 指出来", st["ok"] is False and "hermes" in st["reason"], st["reason"])
        fk.set(ee, "hermes_path", lambda: cli)          # 借用存在的文件当 agent
        st = ee.available()
        check("三件套齐 → ok=True", st["ok"] is True, st["reason"])

        print("[2] 提示合成：文字不改写 + 几何优先墨迹 + 不硬套过远的 size_group")
        blocks = [dict(text="项目概述", role="title", x=0.05, y=0.10, w=0.30, h=0.07, font_size_pt=24),
                  dict(text="正文一", role="body", x=0.05, y=0.40, w=0.30, h=0.05, font_size_pt=15)]
        ink = [dict(box_px=[84, 94, 502, 66], size_group="g1")]
        h = ee.build_hints(blocks, ink, 1672, 941)
        texts = [x["text"] for x in h["lines"]]
        check("文字逐条原样（不改写）", texts == ["项目概述", "正文一"], str(texts))
        check("命中的那块改用墨迹框", h["lines"][0]["box_px"] == [84, 94, 502, 66], str(h["lines"][0]["box_px"]))
        check("命中块拿到 size_group", h["lines"][0]["size_group"] == "g1", str(h["lines"][0]["size_group"]))
        check("远处块不硬套 size_group（正文一 → None）", h["lines"][1]["size_group"] is None,
              str(h["lines"][1]["size_group"]))
        check("无墨迹数据时用归一化坐标换算 px", ee.build_hints(blocks, [], 1672, 941)["lines"][0]["box_px"]
              == [84, 94, 502, 66], str(ee.build_hints(blocks, [], 1672, 941)["lines"][0]["box_px"]))
        check("提示里写明文字不得改写", "不得改写" in h["note"] or "必须与本文件一致" in h["note"], h["note"][:40])

        print("[3] 校验结果判定：不能把缺失/坏数据当成通过")
        page = os.path.join(tmp, "page_001")
        os.makedirs(page, exist_ok=True)
        check("缺 validation.json → 未通过", ee._validation_passed(page)["passed"] is False)
        open(os.path.join(page, "validation.json"), "w").write("{坏 json")
        check("坏 JSON → 未通过且说明原因", ee._validation_passed(page)["passed"] is False)
        json.dump({"passed": False, "errors": ["缺文本"]}, open(os.path.join(page, "validation.json"), "w"))
        check("passed=false → 未通过", ee._validation_passed(page)["passed"] is False)
        json.dump({"passed": True}, open(os.path.join(page, "validation.json"), "w"))
        check("passed=true → 通过", ee._validation_passed(page)["passed"] is True)

        print("[4] 页面目录与 URL 映射")
        for d in ("page_003", "page_001", "page_002"):
            os.makedirs(os.path.join(tmp, "run", "pages", d), exist_ok=True)
        dirs = ee._page_dirs(os.path.join(tmp, "run"))
        check("页面目录按序号排序", [os.path.basename(p) for p in dirs] == ["page_001", "page_002", "page_003"],
              str(dirs))
        check("★ 返回的是绝对路径（曾因只返回目录名导致服务端拼出相对路径而崩）",
              bool(dirs) and all(os.path.isabs(p) and p.endswith(("page_001", "page_002", "page_003")) for p in dirs),
              str(dirs[:1]))
        check("_page_dirs 对不存在的目录返回空", ee._page_dirs(os.path.join(tmp, "nope")) == [])
        fk.set(ee.config, "FILES_DIR", tmp)
        from server.app import pipeline as pl
        fk.set(pl.config, "FILES_DIR", tmp)
        check("路径 → /files/ 相对 URL",
              pl._files_url(os.path.join(tmp, "editppt", "r1", "final", "a.pptx")) == "/files/editppt/r1/final/a.pptx")

        print("[5] convert 早退：不可用 / 无页 / 页缺图（都要给可执行的提示）")
        fk.set(ee, "available", lambda: dict(ok=False, reason="未安装 editppt CLI"))
        r = ee.convert("p_x", [1])
        check("引擎不可用 → ok=False 且带原因", r["ok"] is False and "未安装" in r["error"], r["error"])
        fk.set(ee, "available", lambda: dict(ok=True, reason="可用"))
        fk.set(ee.db, "q", lambda *a, **k: [])
        r = ee.convert("p_x", [1])
        check("没有匹配的页 → 明确报错", r["ok"] is False and "没有找到" in r["error"], r["error"])
        fk.set(ee.db, "q", lambda *a, **k: [dict(page_no=1, image_path="", width=100, height=100,
                                                 native_scene_json="{}")])
        r = ee.convert("p_x", [1])
        check("页缺页面图 → 提示先导入/先生成", r["ok"] is False and "页面图" in r["error"], r["error"])

        print("[6] 失败聚合：任一页没过校验 → 整体失败并保留现场（绝不伪造成功）")
        img = os.path.join(tmp, "p01.png")
        open(img, "wb").write(b"png")
        run_root = os.path.join(tmp, "runs")
        fk.set(ee.config, "FILES_DIR", tmp)
        fk.set(ee.db, "q", lambda *a, **k: [dict(page_no=1, image_path=img, width=100, height=100,
                                                 native_scene_json="{}")])

        def fake_run(args, cwd=None, timeout=0):
            if args and args[0] == "prepare":                      # 模拟 prepare 在 out-root 下自建一层 run 目录
                out_root = args[args.index("--out-root") + 1]
                os.makedirs(os.path.join(out_root, "run1", "pages", "page_001"), exist_ok=True)
                return dict(ok=True, code=0, out="", err="")
            if len(args) >= 2 and args[0] == "run" and args[1] == "status":
                return dict(ok=True, code=0, out=json.dumps({"pages": {}, "active_dispatches": []}), err="")
            if len(args) >= 2 and args[0] == "run" and args[1] == "next":
                return dict(ok=True, code=0, out=json.dumps(
                    {"stage": "dispatch_pages", "suggested_pages": ["page_001"]}), err="")
            return dict(ok=True, code=0, out="", err="")

        fk.set(ee, "_run", fake_run)
        fk.set(ee, "config", ee.config)
        fk.set(ee, "_prepare_page", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("视觉识别失败")))
        fk.set(ee.db, "q", lambda *a, **k: [dict(page_no=1, image_path=img, width=100, height=100,
                                                 native_scene_json="{}")])
        orig_root = ee.config.FILES_DIR
        ee.config.FILES_DIR = tmp
        try:
            r = ee.convert("p_x", [1])
        finally:
            ee.config.FILES_DIR = orig_root
        check("整体 ok=False", r["ok"] is False, json.dumps(r, ensure_ascii=False)[:120])
        check("失败项带页码与阶段", bool(r.get("failures")) and r["failures"][0]["page_no"] == 1,
              json.dumps(r.get("failures"), ensure_ascii=False)[:120])
        check("错误信息给出保留现场路径", "保留现场" in r.get("error", "") or r.get("run_dir"), r.get("error", "")[:80])

        print("[7] 产物发布：复制到 exports/（run 目录只作现场，可随时清理）")
        fk.set(ee.config, "FILES_DIR", tmp)
        src_deck = os.path.join(tmp, "run2", "final", "p01_edited.pptx")
        os.makedirs(os.path.dirname(src_deck), exist_ok=True)
        open(src_deck, "wb").write(b"pptx-bytes")
        dest = ee._publish_deck(src_deck)
        check("发布到 exports/ 并返回新路径", dest.startswith(os.path.join(tmp, "exports")) and os.path.exists(dest),
              dest)
        check("副本内容与源一致", open(dest, "rb").read() == b"pptx-bytes")
        check("★ 发布的是副本（run 目录可随后清理）", dest != src_deck)
        check("复制失败退回原路径（不假装成功）",
              ee._publish_deck(os.path.join(tmp, "nope", "missing.pptx")).endswith("missing.pptx"))

        print("[8] 断点续跑 / 中途可停")
        run = os.path.join(tmp, "resume-run")
        os.makedirs(os.path.join(run, "pages", "page_001"), exist_ok=True)
        os.makedirs(os.path.join(run, "pages", "page_002"), exist_ok=True)
        json.dump({"pages": [{"page_id": "page_001", "status": "recorded"},
                             {"page_id": "page_002", "status": "pending"}]},
                  open(os.path.join(run, "page_jobs.json"), "w"))
        check("_page_status 读到页状态", ee._page_status(run, "page_001") == "recorded")
        check("_page_recorded 只认已记录/已接受",
              ee._page_recorded(run, "page_001") and not ee._page_recorded(run, "page_002"))
        check("_page_status 对未知页返回空串", ee._page_status(run, "page_999") == "")

        fk.set(ee.config, "FILES_DIR", tmp)
        fk.set(ee, "_latest_run_dir", lambda base: base)
        os.makedirs(os.path.join(tmp, "editppt"), exist_ok=True)
        partial = os.path.join(tmp, "editppt", "p_resume-20260921-000000")
        if os.path.exists(partial):
            shutil.rmtree(partial)
        shutil.copytree(run, partial)
        check("★ 有部分完成 → 找到可续跑的 run", ee._latest_incomplete_run("p_resume") == partial, partial)
        json.dump({"pages": [{"page_id": "page_001", "status": "recorded"}]},
                  open(os.path.join(partial, "page_jobs.json"), "w"))
        check("全部完成 → 不再算可续跑", ee._latest_incomplete_run("p_resume") == "")
        st = ee.run_status("p_resume")
        check("run_status 报总数/已完成", st["total"] == 1 and st["recorded"] == 1 and st["has_run"])
        check("全部完成时 resumable=False", st["resumable"] is False, str(st))
        # 服务重启会杀死所有 page worker，一个页面都没来得及 record；现场仍应可续跑（复用 prepare 结果）
        json.dump({"pages": [{"page_id": "page_001", "status": "dispatched"},
                             {"page_id": "page_002", "status": "dispatched"}]},
                  open(os.path.join(partial, "page_jobs.json"), "w"))
        check("★ 0 页完成的中断现场 → 仍找到可续跑的 run",
              ee._latest_incomplete_run("p_resume") == partial)
        st0 = ee.run_status("p_resume")
        check("0 页完成时 resumable=True 且有提示", st0["resumable"] is True and "继续转换" in st0["note"], str(st0))
        json.dump({"pages": [{"page_id": "page_001", "status": "recorded"},
                             {"page_id": "page_002", "status": "pending"}]},
                  open(os.path.join(partial, "page_jobs.json"), "w"))
        json.dump([7, 9], open(os.path.join(partial, "deckforge_pages.json"), "w"))
        check("run_status 带出原始页清单（续跑按它映射）", ee.run_status("p_resume")["page_nos"] == [7, 9],
              str(ee.run_status("p_resume")["page_nos"]))
        check("_stored_page_nos 容错（没写过的现场/空路径）",
              ee._stored_page_nos(run) == [] and ee._stored_page_nos("") == [])
        shutil.rmtree(partial)
        check("没有转换记录时 has_run=False", ee.run_status("p_none")["has_run"] is False)

        print("[8b] run 目录取最新：按 mtime，不按随机 uid 名排序")
        named_big = os.path.join(tmp, "editppt", "p_order-run_zzzzzz")
        named_small = os.path.join(tmp, "editppt", "p_order-run_aaaaaa")
        for path in (named_big, named_small):
            os.makedirs(path, exist_ok=True)
            json.dump({"pages": [{"page_id": "page_001", "status": "pending"}]},
                      open(os.path.join(path, "page_jobs.json"), "w"))
        os.utime(named_big, (1000, 1000))
        os.utime(named_small, (2000, 2000))                    # 名字小但更晚创建
        check("★ 按 mtime 取最新 run（随机 uid 名不能按字典序）",
              ee._latest_run("p_order") == named_small, ee._latest_run("p_order"))
        check("★ 续跑也取 mtime 最新的未完成 run",
              ee._latest_incomplete_run("p_order") == named_small, ee._latest_incomplete_run("p_order"))

        ee.clear_cancel("p_c")
        check("取消令牌默认未置位", ee.is_cancelled("p_c") is False)
        ee.request_cancel("p_c")
        check("request_cancel 后置位", ee.is_cancelled("p_c") is True)
        ee.clear_cancel("p_c")
        check("clear_cancel 后可复位（新作业不受旧取消影响）", ee.is_cancelled("p_c") is False)

        print("[9] 续跑时**跳过已记录页**（不重新派 worker）")
        calls = {"worker": 0}

        def fake_worker(page_dir, page_id, cancel_key=""):
            calls["worker"] += 1
            return dict(ok=True, tail="")

        fk.set(ee, "_run_page_worker", fake_worker)
        fk.set(ee, "_prepare_page", lambda *a, **k: None)
        def fake_run_recorded(args, cwd=None, timeout=0):
            if len(args) >= 2 and args[0] == "run" and args[1] == "status":
                return dict(ok=True, code=0, out=json.dumps(
                    {"pages": {"recorded": ["page_001"]}, "active_dispatches": []}), err="")
            if len(args) >= 2 and args[0] == "run" and args[1] == "next":
                return dict(ok=True, code=0, out=json.dumps({"stage": "finalize"}), err="")
            return dict(ok=True, code=0, out="", err="")
        fk.set(ee, "_run", fake_run_recorded)
        fk.set(ee.db, "q", lambda *a, **k: [dict(page_no=1, image_path=img, width=100, height=100,
                                                 native_scene_json="{}")])
        fk.set(ee, "_publish_deck", lambda deck: deck)
        run2 = os.path.join(tmp, "editppt", "p_res2-20260921-000000")
        os.makedirs(os.path.join(run2, "pages", "page_001"), exist_ok=True)
        os.makedirs(os.path.join(run2, "final"), exist_ok=True)
        open(os.path.join(run2, "final", "x_edited.pptx"), "wb").write(b"pptx")
        json.dump({"pages": [{"page_id": "page_001", "status": "recorded"}]},
                  open(os.path.join(run2, "page_jobs.json"), "w"))
        res2 = ee.convert("p_res2", [1], resume_run_dir=run2)      # 续跑：复用现场、跳过已记录页
        check("已记录页被跳过（worker 未被调用）", calls["worker"] == 0 and res2.get("ok") is True,
              json.dumps(res2, ensure_ascii=False)[:140])
        check("跳过页后被算作已完成", res2.get("pages") == [1])

        print("[10] 状态机与并发上限：run next 解析、并发数收敛")
        fk.set(ee, "_setting", lambda key, default="": "4")
        check("max_concurrent_pages 读设置", ee.max_concurrent_pages() == 4, str(ee.max_concurrent_pages()))
        fk.set(ee, "_setting", lambda key, default="": "99")
        check("并发上限收敛到 8", ee.max_concurrent_pages() == 8, str(ee.max_concurrent_pages()))
        fk.set(ee, "_setting", lambda key, default="": "bad")
        check("非法并发值回落默认 2", ee.max_concurrent_pages() == 2, str(ee.max_concurrent_pages()))

        fk.set(ee, "_run", lambda args, cwd=None, timeout=0: dict(
            ok=True, code=0, out=json.dumps({"stage": "finalize"}), err=""))
        check("_next_json 解析 run next 输出", ee._next_json(tmp).get("stage") == "finalize")
        fk.set(ee, "_run", lambda args, cwd=None, timeout=0: dict(ok=True, code=0, out="not-json", err=""))
        check("run next 输出坏 JSON → stage=error 且带原因",
              ee._next_json(tmp).get("stage") == "error" and ee._next_json(tmp).get("reason"))

        print("[11] 收尾补全：page_result.json 缺失但其它产物齐全时补路径清单")
        pr = os.path.join(tmp, "page_result_case")
        os.makedirs(pr, exist_ok=True)
        for name in ("manifest.json", "imagegen-jobs.json", "page.pptx", "preview.png",
                     "split_assets_contact.png", "validation.json"):
            open(os.path.join(pr, name), "wb").write(b"x")
        check("缺 page_result → 自动补写", ee._ensure_page_result(pr) and os.path.isfile(os.path.join(pr, "page_result.json")))
        payload = json.load(open(os.path.join(pr, "page_result.json"), encoding="utf-8"))
        check("补写内容只是路径清单", payload.get("page_pptx") == "page.pptx" and payload.get("validation") == "validation.json")
        os.remove(os.path.join(pr, "preview.png"))
        os.remove(os.path.join(pr, "page_result.json"))
        check("其它产物缺失 → 不补、不伪造", ee._ensure_page_result(pr) is False)

        print("[12] 整页图校验：imagegen 背景允许，source.png/user-provided 冒充可编辑要拒绝")
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches
        case = os.path.join(tmp, "full_slide_case")
        page = os.path.join(case, "pages", "page_001")
        os.makedirs(page, exist_ok=True)
        img = os.path.join(page, "bg.png")
        Image.new("RGB", (1600, 900), "#071A33").save(img)
        deck = Presentation()
        deck.slide_width = Inches(13.333)
        deck.slide_height = Inches(7.5)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(img, 0, 0, width=deck.slide_width, height=deck.slide_height)
        slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(1)).text = "标题"
        deck_path = os.path.join(case, "deck.pptx")
        deck.save(deck_path)

        def write_manifest(source_type):
            manifest = {"source": {"width_px": 1600, "height_px": 900},
                        "images": [{"path": "bg.png", "box_px": [0, 0, 1600, 900]}],
                        "asset_provenance": [{"path": "bg.png", "source_type": source_type}]}
            json.dump(manifest, open(os.path.join(page, "manifest.json"), "w", encoding="utf-8"))

        write_manifest("imagegen")
        check("imagegen clean base 整页背景允许（背景层）",
              ee._validate_deck_text(deck_path, [dict(page_no=1, truth_texts=["标题"])], case)["ok"] is True)
        write_manifest("user-provided")
        check("user-provided 整页图冒充可编辑 → 拒绝",
              ee._validate_deck_text(deck_path, [dict(page_no=1, truth_texts=["标题"])], case)["ok"] is False)

        print("[13] 取消后不再往下跑（保留现场，不当失败）")
        fk.set(ee, "_run_page_worker", lambda *a, **k: dict(ok=False, cancelled=True, tail="已停止"))
        run3 = os.path.join(tmp, "editppt", "p_res3-20260921-000000")
        os.makedirs(os.path.join(run3, "pages", "page_001"), exist_ok=True)
        json.dump({"pages": [{"page_id": "page_001", "status": "pending"}]},
                  open(os.path.join(run3, "page_jobs.json"), "w"))
        ee.request_cancel("p_res3")
        res3 = ee.convert("p_res3", [1], resume_run_dir=run3, cancel_key="p_res3")
        ee.clear_cancel("p_res3")
        check("取消 → ok=False 且标记 cancelled", res3.get("ok") is False and res3.get("cancelled") is True,
              str(res3.get("error"))[:80])
        check("取消提示保留现场可继续", "可" in str(res3.get("error")) and bool(res3.get("run_dir")))
        check("取消时不做 finalize（partial=True）", res3.get("partial") is True)

        print("[14] 多机派发：辅助机配置 / 槽位顺序 / prompt 路径改写")
        fk.set(ee, "_setting", lambda key, default="": (
            json.dumps([{"host": "root@192.168.110.13", "slots": 2,
                         "workdir": "/opt/editppt-workers",
                         "skill_dir": "/root/.agents/skills/image-to-editable-ppt",
                         "editppt_cli": "/root/.editppt-venv/bin/editppt"}])
            if key == "editppt.remote_workers" else default))
        fk.set(ee, "harness_profile", lambda: "headless")
        fk.set(ee, "agents_skill_dir", lambda: "/platform/agents-skill")
        fk.set(ee, "skill_dir", lambda: "/platform/skill")
        fk.set(ee, "max_concurrent_pages", lambda: 2)
        hubs = ee.remote_workers()
        check("解析 remote worker（host/slots/默认 profile）",
              len(hubs) == 1 and hubs[0]["host"] == "root@192.168.110.13" and hubs[0]["slots"] == 2
              and hubs[0]["profile"] == "headless", str(hubs))
        slots = ee.worker_slots()
        check("★ 槽位 = 本机 2 + 辅助机 2，且本机/远端交替派发",
              len(slots) == 4 and [s["kind"] for s in slots] == ["local", "remote", "local", "remote"],
              str([s["kind"] for s in slots]))
        fk.set(ee, "_setting", lambda key, default="": "not-json" if key == "editppt.remote_workers" else default)
        check("remote_workers 坏 JSON → 空列表（不抛异常）", ee.remote_workers() == [])
        remote_page = "/opt/editppt-workers/run_ab/20261020-1/pages/page_001"
        text = ("runroot=/data/run_ab/20261020-1 page=/data/run_ab/20261020-1/pages/page_001 "
                "worker=/platform/skill agents=/platform/agents-skill "
                "truth=/data/run_ab/20261020-1/pages/page_001/truth_text.txt")
        rewritten = ee._rewrite_remote_prompt(
            text, "/data/run_ab/20261020-1", "/data/run_ab/20261020-1/pages/page_001", remote_page, hubs[0])
        check("★ prompt 里页面/run/技能路径都改写为辅助机路径",
              "runroot=/opt/editppt-workers/run_ab/20261020-1 " in rewritten
              and remote_page in rewritten
              and "/root/.agents/skills/image-to-editable-ppt" in rewritten
              and "/data/run_ab" not in rewritten, rewritten[:200])
        fk.set(ee, "agent_kind", lambda: "harness")
        page_dir = os.path.join(tmp, "run_ab", "20261020-1", "pages", "page_001")
        os.makedirs(page_dir, exist_ok=True)
        prompt = ee._worker_prompt("/data/run_ab/20261020-1", "page_001", page_dir)
        check("★ 平台包装提示含 skill 加载 / 文字真源 / token 纪律（远端也必须用它）",
              "skill" in prompt and "token 预算纪律" in prompt and "worker-prompt.md" in prompt, prompt[:120])
    finally:
        fk.restore()
        shutil.rmtree(tmp, ignore_errors=True)
    print("通过 %d 项，失败 %d 项" % (len(PASS), len(FAIL)))
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
