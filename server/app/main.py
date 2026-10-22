# -*- coding: utf-8 -*-
"""SlideCraft HTTP 层（FastAPI，:8640）。路由严格对齐 docs/API.md。"""
import threading
import asyncio
import json
import os
import re
import shutil
from datetime import datetime
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse

from fastapi import Body, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles

from . import operations, export_snapshots
from . import (beautify, config, db, editppt_engine, page_import, pipeline, provider_usage, providers,
               slide_notes, source_pages, style_builder, stylepacks)
from .paths import app_version

# 版本号唯一真源是仓库根 `VERSION`（别在这里写死字面量；门禁 tools/check_version_sync.py 会拦）。
VERSION = app_version()

app = FastAPI(title="SlideCraft", version=VERSION)
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

os.makedirs(config.FILES_DIR, exist_ok=True)


@app.on_event("startup")
def _startup() -> None:
    db.init_db()
    pipeline.recover_interrupted_jobs()
    pipeline.slide_versions.backfill()


# ────────────────────────── 健康与设置 ──────────────────────────

@app.get("/api/health")
def health() -> Dict[str, Any]:
    caps = providers.capabilities()
    return dict(ok=True, version=VERSION, port=int(os.environ.get("DECKFORGE_PORT", "8640")), db="ok",
                providers=dict(image=caps["image"], text=caps["text"], picked=caps["picked"]),
                network=dict(mode=config.network_mode(), http_proxy=config.http_proxy_url(),
                             domestic_bypass_proxy=True))


@app.get("/api/settings")
def get_settings() -> Dict[str, Any]:
    return dict(providers=config.image_providers(), text=config.text_provider(), vision=config.vision_provider(),
                network=dict(mode=config.network_mode(), http_proxy=config.http_proxy_url()),
                priority=config.priority())


@app.put("/api/settings")
def put_settings(payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    if "provider_id" in payload:
        pid = payload["provider_id"]
        if payload.get("api_key"):
            config.set_setting("key:%s" % pid, payload["api_key"])
        if payload.get("default_model"):
            config.set_setting("model:%s" % pid, payload["default_model"])
        if payload.get("default_size"):
            config.set_setting("size:%s" % pid, payload["default_size"])
        if payload.get("base_url"):
            base_url = str(payload["base_url"]).rstrip("/")
            parsed = urlparse(base_url)
            if parsed.scheme != "https" or not parsed.netloc:
                raise HTTPException(400, "API Base URL 必须是完整的 HTTPS 地址，例如 https://api.guaihub.cc/v1")
            config.set_setting("base_url:%s" % pid, base_url)
        if "unit_price_cny" in payload:
            # 单价由用户填（我们不猜）：空字符串 = 清除覆盖，回落 yaml（可能仍是未知）
            raw_price = str(payload.get("unit_price_cny") or "").strip()
            if not raw_price:
                config.set_setting("price:%s" % pid, "")
            else:
                try:
                    value = float(raw_price)
                except ValueError:
                    raise HTTPException(400, "单价必须是数字（元/张）") from None
                if value < 0:
                    raise HTTPException(400, "单价不能为负数")
                config.set_setting("price:%s" % pid, "%g" % value)
    if payload.get("text_model"):
        config.set_setting("model:deepseek", payload["text_model"])
    net = payload.get("network") or {}
    if net.get("mode"):
        config.set_setting("network:mode", net["mode"])
    if net.get("http_proxy"):
        config.set_setting("network:http_proxy", net["http_proxy"])
    return get_settings()                            # 热加载：无需重启


@app.post("/api/settings/test")
def test_settings(payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    pid = payload.get("provider_id", "")
    if not isinstance(pid, str):
        raise HTTPException(400, "provider_id 必须是字符串")
    api_key = payload.get("api_key")
    if api_key is not None and not isinstance(api_key, str):
        raise HTTPException(400, "api_key 必须是字符串")
    return providers.test_provider(pid, api_key)


# ────────────────────────── 风格包 ──────────────────────────

@app.get("/api/style-packs")
def list_style_packs() -> Dict[str, Any]:
    return dict(items=stylepacks.list_style_packs())


@app.get("/api/style-packs/{sid}")
def get_style_pack(sid: str) -> Dict[str, Any]:
    sp = stylepacks.get_style_pack(sid)
    if not sp:
        raise HTTPException(404, "风格包不存在：%s" % sid)
    return sp


@app.get("/api/layouts")
def list_layouts() -> Dict[str, Any]:
    return dict(items=[{k: v for k, v in it.items() if k != "body"} for it in stylepacks.list_layouts()])


# ────────────────────────── 项目 ──────────────────────────

def project_cover_url(pid: str) -> str:
    """项目封面：**只取成品类图片**（已生成结果 / 美化结果 / 导入成品页），没有就返回空串。

    原稿缩略图（source-pages/）**不算封面**——那只是原稿，拿它当封面会让人以为已经生成。
    """
    row = db.q1("SELECT image_path FROM slides WHERE project_id=? AND image_path IS NOT NULL AND image_path<>'' "
                "ORDER BY COALESCE(display_order, page_no) LIMIT 1", (pid,))
    if row and os.path.isfile(row["image_path"] or ""):
        return pipeline._files_url(row["image_path"])
    for sub in ("beautify/beauty", "pages"):
        folder = os.path.join(config.FILES_DIR, pid, *sub.split("/"))
        if os.path.isdir(folder):
            names = sorted(name for name in os.listdir(folder) if name.lower().endswith(".png"))
            if names:
                return "/files/%s/%s/%s" % (pid, sub, names[0])
    return ""


def project_export_briefs(pid: str) -> List[Dict[str, Any]]:
    """项目导出记录（列表页「下载」用）：只带下载必需字段，不带快照明细。"""
    items = []
    for row in db.q('SELECT * FROM exports WHERE project_id=? ORDER BY created_at DESC', (pid,)):
        full = export_snapshots.response(dict(row))
        items.append(dict(id=full.get("id"), kind=full.get("kind"), file=full.get("file") or "",
                          pages=full.get("pages") or 0, size_bytes=full.get("size_bytes") or 0,
                          created_at=full.get("created_at")))
    return items


@app.get("/api/projects")
def list_projects() -> Dict[str, Any]:
    rows = db.q("SELECT p.*, (SELECT COUNT(*) FROM slides s WHERE s.project_id=p.id) AS page_count "
                "FROM projects p ORDER BY datetime(p.created_at) DESC")
    return dict(items=[dict(get_project(r["id"]), page_count=r["page_count"],
                            cover_url=project_cover_url(r["id"]),
                            exports=project_export_briefs(r["id"])) for r in rows])


@app.post("/api/projects")
def create_project(payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    pid = db.uid("p")
    db.insert("projects", dict(id=pid, title=payload.get("title") or "未命名项目",
                               source_type=payload.get("source_type", "markdown"),
                               intake_mode=payload.get("intake_mode", "doc"),
                               status="draft", created_at=db.now(), updated_at=db.now()))
    return dict(id=pid)


CONFIG_LOCK = threading.RLock()


def guard_generation(pid, page_nos=None):
    for row in db.q("SELECT id FROM jobs WHERE project_id=? AND kind='generate' AND status IN ('queued','running','cancelling')", (pid,)):
        job = pipeline.get_job(row["id"]) or {}
        pages = json.loads(job.get("pages_json") or "{}")
        locked = {int(no) for no, event in pages.items() if event.get("status") in ("queued", "running")}
        if page_nos is None or not pages or locked.intersection(page_nos):
            raise HTTPException(409, "页面任务运行中，暂不能修改内容、配置或重复提交；可查看已有结果")


@app.get("/api/projects/{pid}/jobs")
def project_jobs(pid: str):
    return dict(items=[pipeline.get_job(r["id"]) for r in db.q(
        "SELECT id FROM jobs WHERE project_id=? ORDER BY rowid DESC", (pid,))])


def project_flow(pid: str) -> str:
    """项目属于哪条流程：`beautify`（PPT 美化）/ `editable`（图片转可编辑）/ `wizard`（文档生成向导）。

    判据都是磁盘事实：`data/files/<pid>/beautify/` 由美化上传落地；`data/files/<pid>/import/`
    由「图片转可编辑」导入原页时落地；两者都没有即视为文档生成向导。
    这里只做事实判断，不猜测「用户想用哪条流程」。
    """
    root = os.path.join(config.FILES_DIR, pid)
    if os.path.isdir(os.path.join(root, "beautify")):
        return "beautify"
    if os.path.isdir(os.path.join(root, "import")):
        return "editable"
    return "wizard"


@app.get("/api/projects/{pid}")
def get_project(pid: str) -> Dict[str, Any]:
    row = db.q1("SELECT * FROM projects WHERE id=?", (pid,))
    if not row:
        raise HTTPException(404, "项目不存在")
    data = dict(row)
    data["page_drafts"] = json.loads(data.pop("page_drafts_json", None) or "[]")
    rows = db.q("SELECT * FROM slides WHERE project_id=? ORDER BY page_no", (pid,))
    generated = any(r["image_path"] or r["status"] in ("running", "failed", "interrupted") for r in rows)
    legacy_content = data["content_confirmed_at"] is None and bool(rows) and bool(data["style_pack_id"])
    data["pages_confirmed"] = bool(data["pages_confirmed_at"]) or legacy_content or generated
    data["content_confirmed"] = bool(data["content_confirmed_at"]) or legacy_content
    data["content_revisions"] = {r["id"]: r["content_revision"] for r in rows}
    data["resume_step"] = (6 if generated else 4 if data["content_confirmed"] else
                           3 if data["pages_confirmed"] else 2 if data["parsed_text"] or rows else 1)
    data["flow"] = project_flow(pid)
    return data


@app.patch("/api/projects/{pid}")
def configure_project(pid: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    with CONFIG_LOCK:
        project = get_project(pid)
        if payload.get("expected_updated_at") is not None and payload["expected_updated_at"] != project["updated_at"]:
            raise HTTPException(409, "项目配置已由其他窗口修改，请保留草稿并比较服务器版本")
        if set(payload) - {"title", "expected_updated_at"}:
            guard_generation(pid)
        upd = {}
        if "title" in payload:
            title = payload["title"]
            if not isinstance(title, str) or not title.strip():
                raise HTTPException(400, "项目名称不能为空")
            upd["title"] = title.strip()
        if "style_pack_id" in payload:
            if not stylepacks.get_style_pack(payload["style_pack_id"]):
                raise HTTPException(400, "风格包不存在")
            upd["style_pack_id"] = payload["style_pack_id"]
        if "size" in payload:
            if payload["size"] not in ("1K", "2K", "4K"):
                raise HTTPException(400, "尺寸无效")
            upd["size"] = payload["size"]
        if "target_pages" in payload:
            value = payload["target_pages"]
            if type(value) is not int or not 1 <= value <= 80:
                raise HTTPException(400, "目标页数须为 1–80")
            upd["target_pages"] = value
        if "page_drafts" in payload:
            pages = payload["page_drafts"]
            if not isinstance(pages, list) or not pages:
                raise HTTPException(400, "至少要保留 1 页")
            validate_page_revisions(pid, pages)
            upd.update(page_drafts_json=json.dumps(pages, ensure_ascii=False),
                       target_pages=sum(bool(p.get("keep", True)) for p in pages),
                       pages_confirmed_at="", content_confirmed_at="")
            for p in pages:
                row = db.q1("SELECT * FROM slides WHERE project_id=? AND page_no=?", (pid, p.get("page_no")))
                if row:
                    db.update("slides", row["id"], dict(keep=int(bool(p.get("keep", True)) and bool(row["source_available"])),
                                                       layout=p.get("layout", "cards_4")))
                    if "title" in p and p["title"] != row["title"]:
                        db.update("slides", row["id"], dict(title=p["title"], content_revision=row["content_revision"] + 1,
                            source_content_json=original_content(row)))
        if payload.get("confirm_content"):
            if not project["pages_confirmed"]:
                raise HTTPException(400, "先完成页面确认")
            selected = db.q("SELECT * FROM slides WHERE project_id=? AND keep=1", (pid,))
            if not selected or any(not content_ready(dict(r)) for r in selected):
                raise HTTPException(400, "所选页面内容待补齐")
            upd["content_confirmed_at"] = db.now()
        db.update("projects", pid, upd)
        return get_project(pid)


def _under_files(path: str) -> bool:
    """只允许删 `data/files/` 下的路径。

    DB 里存的是历史绝对路径（导出件、转换现场），一旦指向仓库外（老数据、手工改过），
    跟着删就是灾难；这里只放行 FILES_DIR 之内的目标。
    """
    try:
        root = os.path.realpath(config.FILES_DIR)
        target = os.path.realpath(path)
    except (OSError, TypeError):
        return False
    return target == root or target.startswith(root + os.sep)


def _safe_pid(pid: str) -> bool:
    """pid 会被当成 `data/files/` 下的一段目录名；只认 db.uid 生成的那种形态。"""
    return bool(pid) and re.fullmatch(r"[A-Za-z0-9_\-]+", pid) is not None


def _remove_path(path: str) -> bool:
    if not path or not _under_files(path) or not os.path.exists(path):
        return False
    if os.path.isdir(path):
        shutil.rmtree(path, ignore_errors=True)
    else:
        try:
            os.remove(path)
        except OSError:
            return False
    return not os.path.exists(path)


def _editppt_runs(pid: str) -> List[str]:
    """精细重建的现场目录：`data/files/editppt/<pid>-<runid>/`。"""
    root = os.path.join(config.FILES_DIR, "editppt")
    if not os.path.isdir(root):
        return []
    return [os.path.join(root, name) for name in sorted(os.listdir(root)) if name.startswith(pid + "-")]


def _frozen_inputs(run_dirs: List[str]) -> List[str]:
    """run 冻结的输入图目录（`editppt/input_*`）；路径记在该 run 的 manifest 里。

    必须在删 run 目录**之前**读 manifest，否则这些输入图（拷进来的整页图）会永远留在磁盘上。
    """
    found: List[str] = []
    for run in run_dirs:
        for base, _dirs, names in os.walk(run):
            if "deckforge_snapshot.json" not in names:
                continue
            try:
                with open(os.path.join(base, "deckforge_snapshot.json"), encoding="utf-8") as handle:
                    rows = json.load(handle)
            except (OSError, ValueError):
                continue
            for row in rows if isinstance(rows, list) else []:
                folder = os.path.dirname(str((row or {}).get("image_path") or ""))
                if os.path.basename(folder).startswith("input_") and _under_files(folder) and folder not in found:
                    found.append(folder)
    return found


def _purge_project_files(pid: str, exports: List[Dict[str, Any]]) -> Dict[str, int]:
    """删项目在磁盘上的全部痕迹：项目目录、导出件与其快照目录、精细重建现场。"""
    runs = _editppt_runs(pid)
    frozen = _frozen_inputs(runs)
    removed = dict(project_dir=int(_safe_pid(pid) and _remove_path(os.path.join(config.FILES_DIR, pid))),
                   editppt_runs=sum(1 for path in runs if _remove_path(path)),
                   frozen_inputs=sum(1 for path in frozen if _remove_path(path)),
                   export_files=0, export_dirs=0)
    for row in exports:
        if _remove_path(str(row.get("path") or "")):
            removed["export_files"] += 1
        eid = str(row.get("id") or "")
        # eid 为空时绝不能拼出 `files/exports/` 本身（会把所有项目的导出件一起删掉）
        if eid and _remove_path(os.path.join(config.FILES_DIR, "exports", eid)):
            removed["export_dirs"] += 1
    return removed


@app.delete("/api/projects/{pid}")
def delete_project(pid: str) -> Dict[str, Any]:
    """删除项目：页面/任务/导出记录与其磁盘文件一并清掉。

    - 项目有任务在跑（queued/running/cancelling）或原稿缩略图在后台构建 → **409**，不删半截。
    - 用量账本（`usage_log`）**保留**：成本记录是可审计的账，删项目不等于抹账；
      之后 `GET /api/projects/{pid}/usage` 与全局用量页仍能看到这些记录（项目名显示「已删除项目」）。
    """
    with CONFIG_LOCK:
        row = db.q1("SELECT * FROM projects WHERE id=?", (pid,))
        if not row:
            raise HTTPException(404, "项目不存在")
        live = db.q("SELECT kind FROM jobs WHERE project_id=? AND status IN ('queued','running','cancelling')", (pid,))
        if live:
            kinds = "、".join(sorted({r["kind"] or "任务" for r in live}))
            raise HTTPException(409, "项目还有任务在跑（%s）：请先在工作台取消或等它结束后再删除" % kinds)
        if source_pages.is_running(pid):
            raise HTTPException(409, "原稿缩略图正在后台生成，请等它结束后再删除")

        exports = [dict(r) for r in db.q("SELECT id, path FROM exports WHERE project_id=?", (pid,))]
        conn = db.connect()
        conn.execute("BEGIN IMMEDIATE")
        try:
            # 事务内再查一次：避免「检查完到删除之间」刚好有新任务登记进来。
            if conn.execute("SELECT 1 FROM jobs WHERE project_id=? AND status IN ('queued','running','cancelling') LIMIT 1",
                            (pid,)).fetchone():
                raise HTTPException(409, "项目刚登记了任务，请等它结束后再删除")
            deleted: Dict[str, int] = {}
            deleted["slide_versions"] = conn.execute(
                "DELETE FROM slide_versions WHERE slide_id IN (SELECT id FROM slides WHERE project_id=?)", (pid,)).rowcount
            for table in ("slides", "jobs", "exports", "operation_submissions", "operation_quotes"):
                deleted[table] = conn.execute("DELETE FROM %s WHERE project_id=?" % table, (pid,)).rowcount
            deleted["projects"] = conn.execute("DELETE FROM projects WHERE id=?", (pid,)).rowcount
            conn.execute("COMMIT")
        except BaseException:
            if conn.in_transaction:
                conn.execute("ROLLBACK")
            raise

        for job_id, job in list(pipeline.JOBS.items()):        # 内存作业登记：终态记录也一并清掉
            if job.get("project_id") == pid:
                pipeline._evict_job(job_id)
        files = _purge_project_files(pid, exports)
        kept = int((db.q1("SELECT COUNT(*) AS n FROM usage_log WHERE project_id=?", (pid,)) or {"n": 0})["n"])
        return dict(ok=True, id=pid, title=row["title"], deleted=deleted, files=files, usage_log_kept=kept,
                    note="项目记录与磁盘文件已删除；用量账本保留 %d 条成本记录，删除项目不抹账" % kept)


# 上传上限：前端先挡一道，后端这里再流式校验（DECKFORGE_MAX_UPLOAD_MB 可按机器内存调整）
MAX_UPLOAD_MB = max(1, int(os.environ.get("DECKFORGE_MAX_UPLOAD_MB") or 200))
MAX_UPLOAD_BYTES = MAX_UPLOAD_MB * 1024 * 1024
_UPLOAD_CHUNK = 1024 * 1024


async def _save_upload(file: "UploadFile", path: str, limit: int = MAX_UPLOAD_BYTES) -> int:
    """边收边写盘并计数：超过上限立刻中断并删掉半截文件。

    不用 `await file.read()` 整块读，是为了让 200MB 级别的上传也不会把整份文件
    压在内存里（NAS 虚拟机只有 1.9GB 内存）。
    """
    total = 0
    open_file = open(path, "wb")
    try:
        with open_file as handle:
            while True:
                chunk = await file.read(_UPLOAD_CHUNK)
                if not chunk:
                    break
                total += len(chunk)
                if total > limit:
                    raise HTTPException(413, "文件超过 %dMB 上限（已接收 %.1fMB），请压缩或拆分后再传"
                                        % (limit // 1048576, total / 1048576))
                handle.write(chunk)
    except HTTPException:
        try:
            os.remove(path)
        except OSError:
            pass
        raise
    return total


@app.post("/api/projects/{pid}/upload")
async def upload(pid: str, file: Optional[UploadFile] = File(None),
                 text: Optional[str] = Body(None, embed=True)) -> Dict[str, Any]:
    guard_generation(pid)
    proj = db.q1("SELECT * FROM projects WHERE id=?", (pid,))
    if not proj:
        raise HTTPException(404, "项目不存在")
    src_dir = os.path.join(config.FILES_DIR, pid)
    os.makedirs(src_dir, exist_ok=True)
    if file is not None:
        path = os.path.join(src_dir, "source" + os.path.splitext(file.filename or "x.txt")[1])
        await _save_upload(file, path)
        try:
            parsed = pipeline.parse_file(path, file.filename or "x.txt")
        except Exception as e:                        # noqa: BLE001
            raise HTTPException(400, "解析失败：%s" % e)
        db.update("projects", pid, dict(source_path=path, parsed_text=parsed,
                                        source_type=os.path.splitext(file.filename or "")[1].lstrip(".") or "txt"))
        # 原稿页缩略图是非阻塞辅助能力：只在 PPTX 上传后后台开始；失败绝不影响上传。
        if path.lower().endswith(".pptx"):
            from pptx import Presentation
            original_pages = []
            for no, original in enumerate(Presentation(path).slides, 1):
                texts = [sh.text_frame.text for sh in original.shapes
                         if sh.has_text_frame and sh.text_frame.text.strip()]
                tables = [dict(type="table", title="", unit="", cells=[
                    [cell.text for cell in row.cells] for row in sh.table.rows])
                    for sh in original.shapes if sh.has_table]
                title = original.shapes.title.text if original.shapes.title else (texts[0] if texts else "")
                points = [t for t in texts if t != title]
                content = dict(title=title, outline="", points=points, modules=tables, source_note="")
                valid = bool(len(original.shapes))
                values = dict(source_page_no=no, display_order=no, title=title, outline="",
                    points_json=json.dumps(points, ensure_ascii=False), modules_json=json.dumps(tables, ensure_ascii=False),
                    source_content_json=json.dumps(content, ensure_ascii=False), source_available=int(valid),
                    keep=int(valid), role="content", layout="cards_4")
                existing = db.q1("SELECT * FROM slides WHERE project_id=? AND page_no=?", (pid, no))
                if existing:
                    if any(existing[k] != values[k] for k in ("title", "outline", "points_json", "modules_json")):
                        values["content_revision"] = existing["content_revision"] + 1
                    db.update("slides", existing["id"], values)
                else:
                    db.insert("slides", dict(id=db.uid("s"), project_id=pid, page_no=no, **values,
                                             created_at=db.now(), updated_at=db.now()))
                original_pages.append(dict(page_no=no, title=title, keep=valid, layout="cards_4"))
            db.x("UPDATE slides SET keep=0,source_available=0 WHERE project_id=? AND page_no>?", (pid, len(original_pages)))
            db.update("projects", pid, dict(page_drafts_json=json.dumps(original_pages, ensure_ascii=False),
                target_pages=len(original_pages), pages_confirmed_at="", content_confirmed_at=""))
            source_pages.clear(pid)                   # 重传原稿才使旧缓存失效
            if source_pages.start(pid, path):
                asyncio.create_task(source_pages.build_background(pid, path))
    else:
        parsed = text or ""
        db.update("projects", pid, dict(parsed_text=parsed, source_type="text"))
    char_count = len(parsed)
    suggested = max(3, min(30, round(char_count / 420) or 3))
    return dict(parsed_text=parsed[:20000], char_count=char_count, suggested_pages=suggested, warnings=[])


def _source_pptx(pid: str) -> Dict[str, Any]:
    proj = db.q1("SELECT source_type, source_path FROM projects WHERE id=?", (pid,))
    if not proj:
        raise HTTPException(404, "项目不存在")
    if (proj["source_type"] or "").lower() != "pptx" or not (proj["source_path"] or "").lower().endswith(".pptx"):
        raise HTTPException(400, "该项目没有 PPTX 原稿")
    return dict(proj)


@app.get("/api/projects/{pid}/source-pages")
def get_source_pages(pid: str) -> Dict[str, Any]:
    _source_pptx(pid)
    return source_pages.response(pid)


@app.post("/api/projects/{pid}/source-pages/build")
async def build_source_pages(pid: str) -> Dict[str, Any]:
    proj = _source_pptx(pid)
    if source_pages.start(pid, proj["source_path"]):
        asyncio.create_task(source_pages.build_background(pid, proj["source_path"]))
    return source_pages.response(pid)


@app.get("/api/projects/{pid}/source-pages/{page_no}")
def get_source_page_large(pid: str, page_no: int, w: int = Query(1100)) -> Response:
    proj = _source_pptx(pid)
    try:
        # 只提供约定的大图档；忽略其它宽度，避免这个只读接口被当作无界图像服务。
        data = source_pages.large_page_png(proj["source_path"], page_no, 1100 if w != 1100 else w)
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc
    except Exception as exc:  # 大图同样是辅助预览，不暴露服务端堆栈
        raise HTTPException(503, "原页大图暂不可用：%s" % exc) from exc
    return Response(content=data, media_type="image/png")


def _file_url(pid: str, path: str) -> str:
    """项目目录内的文件 → 对外 `/files/<pid>/...` URL。

    兼容两种布局：生成的 `‹页号›.png` 与导入的 `pages/pNN.png`（原先是硬编码前者）。
    """
    if not path:
        return ""
    root = os.path.join(config.FILES_DIR, pid)
    rel = os.path.relpath(path, root)
    if rel.startswith(".."):
        return ""
    return "/files/%s/%s" % (pid, rel.replace(os.sep, "/"))


# ────────── 图片转可编辑（首页模式②）：PDF / 图片 / pptx → 每页一张图 + 图片版 ──────────

@app.post("/api/projects/{pid}/import-pages")
async def import_pages(pid: str, files: List[UploadFile] = File(...)) -> Dict[str, Any]:
    """把上传的 PDF / 图片 / pptx 逐页渲染，并登记成 slides。

    登记之后既有的两条流水线可直接复用：
      · 「导出图片版」读 `slides.image_path`；
      · 「转可编辑」以该页图做视觉识别 + 原生重建（不嵌整页图）。
    多张图片按上传顺序拼成一册；pptx 复用原稿缩略图链路的 PDF 缓存（LibreOffice 只转一次）。
    """
    guard_generation(pid)
    proj = db.q1("SELECT * FROM projects WHERE id=?", (pid,))
    if not proj:
        raise HTTPException(404, "项目不存在")

    dest_dir = os.path.join(config.FILES_DIR, pid, "import")
    os.makedirs(dest_dir, exist_ok=True)
    saved: List[str] = []
    for index, item in enumerate(files or []):
        ext = os.path.splitext(item.filename or "")[1].lower()
        path = os.path.join(dest_dir, "src%02d%s" % (index + 1, ext or ".bin"))
        written = await _save_upload(item, path)
        if not written:                       # 空文件跳过（与原逻辑一致）
            try:
                os.remove(path)
            except OSError:
                pass
            continue
        saved.append(path)
    if not saved:
        raise HTTPException(400, "没有收到可用的文件")

    try:
        st = await asyncio.to_thread(page_import.render_pages, pid, saved)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:                                    # 转换器故障按降级处理，不吐堆栈
        raise HTTPException(503, "原稿渲染失败：%s" % exc) from exc

    count = page_import.register_slides(pid, st)
    # 导入的页面图同样是"当前版本"资产：立刻登记版本记录。
    # 否则 current_version_id 为空，「转可编辑」会以「转换页面或当前版本记录缺失」拒绝，
    # 而 backfill() 只在服务启动时跑 —— 表现为"新建项目后必须重启才能转换"。
    pipeline.slide_versions.backfill()
    # 页面确认步骤的「原稿缩略图」来自 source_pages；导入的原稿同样要登记，
    # 否则该步骤只会显示「原稿缩略图暂不可用」。该模块任何失败只写状态文件，不影响导入结果。
    if source_pages.start(pid, saved[0]):
        asyncio.create_task(source_pages.build_background(pid, saved[0]))
    db.update("projects", pid, dict(
        source_type=os.path.splitext(saved[0])[1].lstrip(".") or "import", source_path=saved[0],
        status="imported", target_pages=count))
    return dict(
        page_count=count, renderer=st.get("renderer"),
        pages=[dict(page_no=p["page_no"], image_url=_file_url(pid, page_import.page_image(pid, p["page_no"])),
                    width=p["width"], height=p["height"]) for p in st["pages"]],
        image_deck="/api/projects/%s/export?kind=image" % pid,
        editable="/api/projects/%s/editable" % pid)


@app.post("/api/projects/{pid}/outline")
def outline(pid: str, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    guard_generation(pid)
    proj = db.q1("SELECT * FROM projects WHERE id=?", (pid,))
    if not proj:
        raise HTTPException(404, "项目不存在")
    if proj["source_type"] == "pptx" and db.q1("SELECT id FROM slides WHERE project_id=? AND source_page_no IS NOT NULL", (pid,)):
        pages = [dict(page_no=r["page_no"], title=r["title"], role=r["role"], outline=r["outline"],
                      points=json.loads(r["points_json"] or "[]"),
                      chapter_no=r["chapter_no"] or 0, chapter_title=r["chapter_title"] or "")
                 for r in db.q("SELECT * FROM slides WHERE project_id=? ORDER BY page_no", (pid,))]
        return dict(pages=pages, chapters=pipeline.chapter_summary(pages),
                    configured=False, fellback="source_extract", hint="直接采用原稿逐页文字与表格；图片页内容请人工补齐，不重跑已成功页面")
    target = int(payload.get("target_pages") or 0)
    res = pipeline.build_outline(proj["parsed_text"] or "", target)
    pages = res["pages"]
    drafts = {p["page_no"]: p for p in json.loads(proj["page_drafts_json"] or "[]")}
    for p in pages:
        existing = db.q1("SELECT * FROM slides WHERE project_id=? AND page_no=?", (pid, p["page_no"]))
        draft = drafts.get(p["page_no"], {})
        values = dict(title=draft.get("title") or p["title"], role=p.get("role", "content"),
                      layout=draft.get("layout") or p.get("layout", "cards_4"),
                      outline=p.get("outline", ""), points_json=json.dumps(p.get("points", []), ensure_ascii=False),
                      must_keep_json=json.dumps(p.get("must_keep", []), ensure_ascii=False),
                      chapter_no=int(p.get("chapter_no") or 0),
                      chapter_title=str(p.get("chapter_title") or "")[:24],
                      keep=int(bool(draft.get("keep", True))))
        if existing:
            if any(existing[k] != v for k, v in values.items() if k not in ("keep",)):
                if not existing["source_content_json"]:
                    values["source_content_json"] = json.dumps(dict(title=existing["title"], outline=existing["outline"],
                        points=json.loads(existing["points_json"] or "[]")), ensure_ascii=False)
                values["content_revision"] = existing["content_revision"] + 1
            db.update("slides", existing["id"], values)
        else:
            db.insert("slides", dict(id=db.uid("s"), project_id=pid, page_no=p["page_no"],
                                     display_order=p["page_no"], source_content_json=json.dumps(dict(title=values["title"],
                                         outline=values["outline"], points=json.loads(values["points_json"])), ensure_ascii=False), **values,
                                     status="pending", created_at=db.now(), updated_at=db.now()))
    db.x("DELETE FROM slides WHERE project_id=? AND page_no>? AND image_path IS NULL", (pid, len(pages)))
    db.update("projects", pid, dict(status="outlined", target_pages=len(pages),
                                   pages_confirmed_at="", content_confirmed_at=""))
    return dict(pages=pages, chapters=res.get("chapters") or pipeline.chapter_summary(pages),
                configured=res["configured"], fellback=res.get("fellback", ""),
                hint=res.get("hint", ""), provider=res.get("provider"), model=res.get("model"))


def original_content(row) -> str:
    return row["source_content_json"] or json.dumps(dict(title=row["title"], outline=row["outline"],
        points=json.loads(row["points_json"] or "[]"), modules=json.loads(row["modules_json"] or "[]"),
        source_note=row["source_note"]), ensure_ascii=False)


def validate_page_revisions(pid: str, pages: List[Dict[str, Any]]) -> None:
    for page in pages:
        row = db.q1("SELECT * FROM slides WHERE project_id=? AND page_no=?", (pid, page.get("page_no")))
        if row and page.get("title", row["title"]) != row["title"] and "revision" in page and page["revision"] != row["content_revision"]:
            raise HTTPException(409, "页面内容已变更，请保留草稿并重新加载")


@app.put("/api/projects/{pid}/pages")
def put_pages(pid: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    guard_generation(pid)
    get_project(pid)
    pages = payload.get("pages", [])
    if not pages or not any(p.get("keep", True) for p in pages):
        raise HTTPException(400, "至少要保留 1 页")
    validate_page_revisions(pid, pages)
    db.update("projects", pid, dict(page_drafts_json=json.dumps(pages, ensure_ascii=False),
                                   target_pages=sum(bool(p.get("keep", True)) for p in pages),
                                   pages_confirmed_at=db.now(), content_confirmed_at=""))
    for p in pages:
        row = db.q1("SELECT * FROM slides WHERE project_id=? AND page_no=?", (pid, p.get("page_no")))
        if not row:
            continue
        upd = {}
        content_changed = False
        if "keep" in p:
            upd["keep"] = 1 if p["keep"] and row["source_available"] else 0
        if p.get("title") and p["title"] != row["title"]:
            upd["title"] = p["title"]
            content_changed = True
        if "outline" in p and str(p.get("outline") or "") != (row["outline"] or ""):
            upd["outline"] = str(p.get("outline") or "")[:8000]
            content_changed = True
        if "points" in p:
            before_points = json.loads(row["points_json"] or "[]")
            next_points = [str(x) for x in (p.get("points") or [])][:60]
            if next_points != before_points:
                upd["points_json"] = json.dumps(next_points, ensure_ascii=False)
                content_changed = True
        if content_changed:
            # 大纲编辑里改的标题/摘要/原文都算内容变更：统一 bump 一次修订号并留原文快照
            upd["content_revision"] = row["content_revision"] + 1
            upd["source_content_json"] = original_content(row)
        if p.get("layout"):
            upd["layout"] = p["layout"]
        if "chapter_no" in p:
            upd["chapter_no"] = int(p.get("chapter_no") or 0)
        if "chapter_title" in p:
            upd["chapter_title"] = str(p.get("chapter_title") or "")[:24]
        if upd:
            db.update("slides", row["id"], upd)
    return dict(ok=True, content_revisions={r["id"]: r["content_revision"] for r in db.q(
        "SELECT * FROM slides WHERE project_id=?", (pid,))})


def content_ready(row: Dict[str, Any]) -> bool:
    points = row.get("points") or json.loads(row.get("points_json") or "[]")
    modules = row.get("modules") or json.loads(row.get("modules_json") or "[]")
    material = any(m.get("text") or any(c is not None and str(c).strip() for line in m.get("cells", []) for c in line)
                   for m in modules if isinstance(m, dict))
    return bool(str(row.get("title") or "").strip() and
                (str(row.get("outline") or "").strip() or any(str(p).strip() for p in points) or material))


# ── 页面删除 / 撤销恢复（内容梳理里删除单页）──────────────────────────────
# 硬删除 + 快照撤销：不引入 deleted 列，避免所有既有查询都要加过滤。
# 撤销时按快照重建该页（id 会换新，内容/图片/章节等逐字段还原），并放回原页码。
_SLIDE_RESTORE_FIELDS = (
    "title", "role", "layout", "outline", "points_json", "must_keep_json", "keep", "image_path",
    "width", "height", "status", "provider", "model", "size", "cost_cny", "latency_ms", "editable",
    "native_scene_json", "native_status", "native_object_count", "native_warnings_json", "verify_json",
    "error", "source_page_no", "source_available", "content_revision", "generated_revision",
    "modules_json", "source_note", "source_content_json", "material_strategy", "generated_style",
    "generated_content_json", "reviewed_image", "chapter_no", "chapter_title",
)


def _renumber_slides(pid: str) -> None:
    """按当前顺序把页码/展示序号重排为 1..N（删除或恢复后调用）。"""
    rows = db.q("SELECT id, page_no, display_order FROM slides WHERE project_id=? ORDER BY page_no, display_order", (pid,))
    for index, row in enumerate(rows, 1):
        if row["page_no"] != index or row["display_order"] != index:
            db.x("UPDATE slides SET page_no=?, display_order=? WHERE id=?", (index, index, row["id"]))


def _sync_pages_after_mutation(pid: str) -> None:
    """删除/恢复后同步项目侧状态：页数、页清单草稿（按 slides 重建）、内容确认失效。"""
    rows = db.q("SELECT * FROM slides WHERE project_id=? ORDER BY page_no", (pid,))
    drafts = [dict(page_no=row["page_no"], title=row["title"], keep=bool(row["keep"]),
                   layout=row["layout"], chapter_no=row["chapter_no"] or 0,
                   chapter_title=row["chapter_title"] or "") for row in rows]
    db.update("projects", pid, dict(
        target_pages=len(rows),
        page_drafts_json=json.dumps(drafts, ensure_ascii=False),
        pages_confirmed_at=db.now(),
        content_confirmed_at="",
    ))


@app.delete("/api/projects/{pid}/slides/{slide_id}")
def delete_slide(pid: str, slide_id: str) -> Dict[str, Any]:
    """删除单页（内容梳理里用）。返回删除前的完整快照，前端据此提供 10 秒撤销。"""
    guard_generation(pid)
    get_project(pid)
    row = db.q1("SELECT * FROM slides WHERE project_id=? AND id=?", (pid, slide_id))
    if not row:
        raise HTTPException(404, "页面不存在")
    total = int((db.q1("SELECT COUNT(*) AS n FROM slides WHERE project_id=?", (pid,)) or {"n": 0})["n"])
    if total <= 1:
        raise HTTPException(400, "至少要保留 1 页；如需换页请先新增页面")
    snapshot = dict(row)
    db.x("DELETE FROM slides WHERE id=?", (slide_id,))
    db.x("DELETE FROM slide_versions WHERE slide_id=?", (slide_id,))
    _renumber_slides(pid)
    _sync_pages_after_mutation(pid)
    return dict(ok=True, deleted=snapshot, page_no=int(snapshot.get("page_no") or 0))


@app.post("/api/projects/{pid}/slides/restore")
def restore_deleted_slide(pid: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    """撤销删除：按前端回传的快照重建该页，并放回原页码。"""
    guard_generation(pid)
    get_project(pid)
    page = payload.get("page") or {}
    if not isinstance(page, dict) or not page:
        raise HTTPException(400, "缺少要恢复的页面快照")
    try:
        page_no = int(payload.get("page_no") or page.get("page_no") or 0)
    except (TypeError, ValueError):
        page_no = 0
    if page_no < 1:
        raise HTTPException(400, "缺少恢复位置（page_no）")
    existing = db.q("SELECT id, page_no FROM slides WHERE project_id=? ORDER BY page_no", (pid,))
    page_no = min(page_no, len(existing) + 1)
    for row in sorted([r for r in existing if r["page_no"] >= page_no], key=lambda r: -r["page_no"]):
        db.x("UPDATE slides SET page_no=?, display_order=? WHERE id=?",
             (row["page_no"] + 1, row["page_no"] + 1, row["id"]))
    record = dict(id=db.uid("s"), project_id=pid, page_no=page_no, display_order=page_no,
                  created_at=db.now(), updated_at=db.now())
    for field in _SLIDE_RESTORE_FIELDS:
        if field in page:
            record[field] = page[field]
    record.setdefault("title", "恢复的页面")
    record["keep"] = 1
    db.insert("slides", record)
    _renumber_slides(pid)
    _sync_pages_after_mutation(pid)
    return dict(ok=True, slide_id=record["id"], page_no=page_no)


@app.patch("/api/projects/{pid}/slides/{slide_id}/content")
def save_content(pid: str, slide_id: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    page = db.q1("SELECT page_no FROM slides WHERE project_id=? AND id=?", (pid, slide_id))
    guard_generation(pid, [page["page_no"]] if page else [])
    conn = db.connect()
    conn.execute("BEGIN IMMEDIATE")
    try:
        row = conn.execute("SELECT * FROM slides WHERE project_id=? AND id=?", (pid, slide_id)).fetchone()
        if not row:
            raise HTTPException(404, "页面不存在")
        if payload.get("revision") != row["content_revision"]:
            raise HTTPException(409, "内容已由其他窗口修改，请保留草稿并重新加载")
        upd = {}
        for key in ("title", "outline", "source_note"):
            if key in payload:
                if not isinstance(payload[key], str):
                    raise HTTPException(400, "内容须为文字")
                upd[key] = payload[key]
        for key in ("points", "modules"):
            if key in payload:
                if not isinstance(payload[key], list):
                    raise HTTPException(400, "要点与模块须为列表")
                if key == "points" and any(not isinstance(p, str) for p in payload[key]):
                    raise HTTPException(400, "每条要点须为文字")
                if key == "modules":
                    for module in payload[key]:
                        if not isinstance(module, dict):
                            raise HTTPException(400, "模块须为对象")
                        cells = module.get("cells")
                        if cells is not None and (not isinstance(cells, list) or any(
                            not isinstance(line, list) or len(line) != len(cells[0]) or any(
                                c is not None and type(c) not in (str, int, float) for c in line) for line in cells)):
                            raise HTTPException(400, "表格须为完整矩形二维单元格")
                upd[key + "_json"] = json.dumps(payload[key], ensure_ascii=False)
        if "material_strategy" in payload:
            if payload["material_strategy"] != "reference":
                raise HTTPException(400, "当前生成仅支持内容参考，不能保证保留原图")
            upd["material_strategy"] = "reference"
        changed = any((json.loads(row[k] or "[]") != json.loads(v)) if k in ("points_json", "modules_json")
                      else row[k] != v for k, v in upd.items())
        if changed:
            if not row["source_content_json"]:
                upd["source_content_json"] = json.dumps(dict(title=row["title"], outline=row["outline"],
                    points=json.loads(row["points_json"] or "[]"), modules=json.loads(row["modules_json"] or "[]"),
                    source_note=row["source_note"]), ensure_ascii=False)
            upd["content_revision"] = row["content_revision"] + 1
            upd["updated_at"] = db.now()
            conn.execute("UPDATE slides SET " + ", ".join(k + "=?" for k in upd) + " WHERE id=?",
                         tuple(upd.values()) + (slide_id,))
            conn.execute("UPDATE projects SET content_confirmed_at='', updated_at=? WHERE id=?", (db.now(), pid))
        conn.execute("COMMIT")
    except BaseException:
        conn.execute("ROLLBACK")
        raise
    return next(p for p in slides(pid)["items"] if p["slide_id"] == slide_id)


@app.post("/api/projects/{pid}/slides/{slide_id}/outline-retry")
def retry_page_content(pid: str, slide_id: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    row = db.q1("SELECT * FROM slides WHERE project_id=? AND id=?", (pid, slide_id))
    if not row:
        raise HTTPException(404, "页面不存在")
    source = json.loads(row["source_content_json"] or "null") or dict(
        title=row["title"], outline=row["outline"], points=json.loads(row["points_json"] or "[]"))
    text = "\n".join(str(v) for v in [source.get("title") or "", source.get("outline") or "", *(source.get("points") or [])] if v)
    if not text.strip():
        raise HTTPException(400, "此页没有可梳理的文字真源，请人工补齐或排除；不会重跑其他页面")
    if payload.get("revision") != row["content_revision"]:
        raise HTTPException(409, "内容修订已变更，请先保存并刷新")
    result = pipeline.build_outline(text, 1)
    if not result.get("pages"):
        raise HTTPException(400, "此页梳理失败，原内容与其他页保持不变")
    page = result["pages"][0]
    return save_content(pid, slide_id, dict(revision=row["content_revision"], title=page.get("title") or "",
                       outline=page.get("outline") or "", points=page.get("points") or []))


@app.post("/api/projects/{pid}/slides/{slide_id}/review")
def review_content(pid: str, slide_id: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    row = db.q1("SELECT * FROM slides WHERE project_id=? AND id=?", (pid, slide_id))
    if not row or not row["image_path"]:
        raise HTTPException(400, "尚无结果可核对")
    if payload.get("image_url") != _file_url(pid, row["image_path"]):
        raise HTTPException(409, "图片已更新，请重新核对")
    db.update("slides", slide_id, dict(reviewed_image=row["image_path"]))
    return dict(ok=True)


def check_generation(pid: str, page_nos: List[int], revisions=None) -> None:
    project = get_project(pid)
    rows = db.q("SELECT * FROM slides WHERE project_id=? AND keep=1", (pid,))
    chosen = [r for r in rows if r["page_no"] in page_nos]
    if len(chosen) != len(set(page_nos)) or not chosen or any(
            not r["source_available"] or not content_ready(dict(r)) for r in chosen):
        raise HTTPException(400, "所选页面内容待补齐或已排除")
    # 旧项目接口兼容；M2 向导必须人工确认。
    if project["content_confirmed_at"] == "":
        raise HTTPException(400, "请先保存并人工确认内容")
    if revisions is not None and any(revisions.get(r["id"]) != r["content_revision"] for r in chosen):
        raise HTTPException(409, "提交修订与已保存内容不一致，请重新确认")


@app.post("/api/projects/{pid}/generate")
@operations.confirmed("generate")
async def generate(pid: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    proj = db.q1("SELECT * FROM projects WHERE id=?", (pid,))
    if not proj:
        raise HTTPException(404, "项目不存在")
    style_id = payload.get("style_pack_id") or "tech-deepblue"
    size = payload.get("size") or "2K"
    rows = db.q("SELECT page_no FROM slides WHERE project_id=? AND keep=1 ORDER BY page_no", (pid,))
    page_nos = payload.get("page_nos") or [r["page_no"] for r in rows]
    if not page_nos:
        raise HTTPException(400, "没有可生成的页面（先在②页面确认里勾选）")
    if any(not content_ready(dict(r)) for r in db.q("SELECT * FROM slides WHERE project_id=? AND keep=1", (pid,))):
        raise HTTPException(400, "全套生成要求所有所选页面内容就绪，请补齐或排除")
    guard_generation(pid, page_nos)
    check_generation(pid, page_nos, payload.get("content_revisions"))
    job_id = await pipeline.run_generation(pid, list(page_nos), style_id, size)
    picked = providers.pick_image_provider()
    return dict(job_id=job_id, configured=bool(picked),
                fellback="" if picked else "local_render",
                hint="" if picked else "未配置生图 key：本次用本地兜底渲染出图（结构与导出链路完全可用）",
                provider=(picked or {}).get("id", "local_render"),
                progress_url="/api/jobs/%s/stream" % job_id)


@app.post("/api/projects/{pid}/slides/{slide_id}/versions/{version_id}/restore")
def restore_slide_version(pid: str, slide_id: str, version_id: str) -> Dict[str, Any]:
    with CONFIG_LOCK:
        row = db.q1("SELECT * FROM slides WHERE id=? AND project_id=?", (slide_id, pid))
        if not row:
            raise HTTPException(404, '页面不存在')
        guard_generation(pid, [row['page_no']])
        try:
            pipeline.slide_versions.select(slide_id, version_id)
            pipeline.refresh_project_result_status(pid)
        except ValueError as exc:
            raise HTTPException(400, str(exc))
        return dict(ok=True, current_version_id=version_id)


@app.get("/api/projects/{pid}/slides")
def slides(pid: str) -> Dict[str, Any]:
    rows = db.q("SELECT * FROM slides WHERE project_id=? ORDER BY page_no", (pid,))
    out = []
    for r in rows:
        d = dict(r)
        d["slide_id"] = r["id"]
        d["versions"] = []
        for index, v in enumerate(db.q("SELECT * FROM slide_versions WHERE slide_id=? ORDER BY rowid", (r['id'],)), 1):
            snapshot = json.loads(v['snapshot_json'])
            d['versions'].append(dict(id=v['id'], number=index, current=v['id'] == r['current_version_id'],
                created_at=v['created_at'], generated_revision=snapshot.get('generated_revision'),
                image_url=pipeline._files_url(snapshot['image_path']) if os.path.isfile(snapshot.get('image_path') or '') else ''))
        d["display_order"] = r["display_order"] or r["page_no"]
        d["points"] = json.loads(r["points_json"] or "[]")
        d["modules"] = json.loads(r["modules_json"] or "[]")
        d["source_content"] = json.loads(r["source_content_json"] or "null")
        d["generated_content"] = json.loads(r["generated_content_json"] or "null")
        d["content_ready"] = content_ready(d)
        d["content_stale"] = bool(r["image_path"] and r["provider"] != "import" and
                                   (r["generated_revision"] if r["generated_revision"] is not None else 1) != r["content_revision"])
        project = db.q1("SELECT * FROM projects WHERE id=?", (pid,))
        d["config_stale"] = bool(r["image_path"] and r["generated_style"] and
                                (r["generated_style"] != project["style_pack_id"] or r["size"] != project["size"]))
        d["review_required"] = bool(r["image_path"] and r["provider"] != "import" and r["reviewed_image"] != r["image_path"])
        d["must_keep"] = json.loads(r["must_keep_json"] or "[]")
        # 未生成页没有 verify_json；API 仍返回完整的校验形状，避免前端缩略图读取 issues 时白屏。
        verify = json.loads(r["verify_json"] or "{}")
        d["verify"] = {
            "text_accuracy": verify.get("text_accuracy", 0.0) if isinstance(verify, dict) else 0.0,
            "issues": verify.get("issues", []) if isinstance(verify, dict) else [],
            "scope": "source_only", "image_verified": False,
        }
        d["native_scene"] = json.loads(r["native_scene_json"] or "{}")
        d["native_warnings"] = json.loads(r["native_warnings_json"] or "[]")
        # version=1 的通用布局以及图片更新后的旧场景，都不能再被 UI 称作“原生可编辑”；
        # 但"从未构建(pending)/构建失败(failed)/构建中(building)"也不能被报成"待重建"。
        d["native_status"] = pipeline.reported_native_status(d)
        d["native_object_count"] = int(r["native_object_count"] or 0)
        for task_row in db.q("SELECT id FROM jobs WHERE project_id=? AND kind='generate' AND status IN ('queued','running','cancelling') ORDER BY rowid DESC", (pid,)):
            task = pipeline.get_job(task_row["id"]) or {}
            event = json.loads(task.get("pages_json") or "{}").get(str(r["page_no"]), {})
            if event.get("status") in ("queued", "running"):
                d["status"] = event["status"]
                break
        d["image_url"] = _file_url(pid, r["image_path"]) if r["image_path"] and os.path.isfile(r["image_path"]) and os.path.getsize(r["image_path"]) else ""
        d["task_state"] = "succeeded" if d["status"] == "done" and d["image_url"] else "failed" if d["status"] == "done" else "queued" if d["status"] == "pending" else d["status"]
        d["result_check"] = "needs_attention" if d["content_stale"] or d["config_stale"] or (r["image_path"] and not d["image_url"]) else "unreviewed" if d["review_required"] or not d["image_url"] else "verified"
        d["cost_cny"] = float(r["cost_cny"] or 0.0)
        out.append(d)
    return dict(items=out)


@app.post("/api/projects/{pid}/slides/{page_no}/regenerate")
@operations.confirmed("regenerate")
async def regenerate(pid: str, page_no: int, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    proj = db.q1("SELECT * FROM projects WHERE id=?", (pid,))
    if not proj:
        raise HTTPException(404, "项目不存在")
    guard_generation(pid, [page_no])
    check_generation(pid, [page_no], payload.get("content_revisions"))
    job_id = await pipeline.run_generation(pid, [page_no], proj["style_pack_id"] or "tech-deepblue",
                                          proj["size"] or "2K")
    return dict(job_id=job_id)


@app.get("/api/editable-engines")
def editable_engines() -> Dict[str, Any]:
    """「图片转可编辑」只有一条路：第三方 skill 的精细重建。

    2026-09-21 按要求下线内置引擎（一次视觉识别 + 本地渲染）：那条路产出的「原生可编辑」
    与 skill 的精细重建质量差距过大，混在一起让交互也乱。现在引擎不可用就如实报错，不回退。
    """
    hi = editppt_engine.available()
    return dict(engines=[
        dict(id="editppt", name="精细重建引擎（第三方 skill）", ok=bool(hi.get("ok")),
             reason=hi.get("reason", ""), seconds_per_page=780, cost_note="每页一次 agent 重建（本机模型额度）",
             cli=hi.get("cli", ""), skill_dir=hi.get("skill_dir", ""),
             note="驱动 image-to-editable-ppt skill 的状态机：逐页 agent 重建 + 自带 validation 校验；"
                  "**实测约 13 分钟/页**（1 页 786 秒），产物仅含所转页面"),
    ])


@app.post("/api/projects/{pid}/editable")
@operations.confirmed("editable")
async def to_editable(pid: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    if not db.q1("SELECT id FROM projects WHERE id=?", (pid,)):
        raise HTTPException(404, "项目不存在")
    page_nos = payload.get("page_nos") or []
    if not page_nos:
        raise HTTPException(400, "请先勾选要转可编辑的页面")
    engine = str(payload.get("engine") or "editppt")
    if engine != "editppt":
        raise HTTPException(400, "engine 只支持 editppt：内置引擎（本地渲染）已下线，"
                                 "「图片转可编辑」统一走第三方 skill 的精细重建")
    state = editppt_engine.available()
    if not state.get("ok"):
        raise HTTPException(400, "精细重建引擎不可用：%s" % state.get("reason"))
    resume = bool(payload.get("resume"))
    try:
        job_id = await pipeline.run_editppt(pid, list(page_nos), resume=resume)
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    return dict(job_id=job_id, engine="editppt", pages=len(page_nos), resume=resume,
                estimate=dict(pages=len(page_nos), seconds=len(page_nos) * 780, cost_cny=None),
                mode="editppt_skill",
                message=("继续上次未完成的转换：已完成的页会跳过" if resume else
                         "将由本机 agent 逐页重建（实测约 13 分钟/页），产物自带 validation 校验；"
                         "仅包含所选页面；可随时「停止」，已完成页与现场都保留，之后可继续"))


@app.post("/api/projects/{pid}/editable/cancel")
def cancel_editable(pid: str) -> Dict[str, Any]:
    """停止正在跑的高精度转换：终止当前 page worker，**保留已完成页与现场**（之后可继续）。"""
    if not db.q1("SELECT id FROM projects WHERE id=?", (pid,)):
        raise HTTPException(404, "项目不存在")
    editppt_engine.request_cancel(pid)
    return dict(ok=True, message="已请求停止：当前页会被终止，已完成的页与现场全部保留，可点「继续转换」接着跑")


@app.get("/api/projects/{pid}/editppt")
def editppt_status(pid: str) -> Dict[str, Any]:
    """高精度转换的现场状态：界面据此显示「已完成 N/M 页，可继续」。"""
    if not db.q1("SELECT id FROM projects WHERE id=?", (pid,)):
        raise HTTPException(404, "项目不存在")
    return editppt_engine.run_status(pid)


@app.post("/api/projects/{pid}/export-preview")
def export_preview(pid: str, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    get_project(pid)
    try:
        items = slides(pid)['items']
        for task in db.q("SELECT id FROM jobs WHERE project_id=? AND kind='generate' AND status IN ('queued','running','cancelling')", (pid,)):
            if not json.loads((pipeline.get_job(task['id']) or {}).get('pages_json') or '{}'):
                for item in items:
                    if item['keep']: item['status'] = 'running'
        return export_snapshots.plan(pid, items, payload.get('kind', 'image'), payload.get('page_nos'))
    except ValueError as exc:
        raise HTTPException(400, str(exc))


@app.post("/api/projects/{pid}/exports")
def create_export(pid: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    with CONFIG_LOCK:
        preview = export_preview(pid, payload)
        if payload.get('snapshot_token') != preview['snapshot_token']:
            raise HTTPException(409, '范围或当前版本已变化，请重新确认导出快照')
        try:
            return export_snapshots.create(pid, preview)
        except ValueError as exc:
            raise HTTPException(400, str(exc))


@app.get("/api/projects/{pid}/exports")
def export_history(pid: str) -> Dict[str, Any]:
    get_project(pid)
    return dict(items=[export_snapshots.response(dict(r)) for r in db.q(
        'SELECT * FROM exports WHERE project_id=? ORDER BY created_at DESC', (pid,))])


@app.get("/api/projects/{pid}/export")
def export(pid: str, kind: str = Query("image")) -> Dict[str, Any]:
    if kind == "image":
        guard_generation(pid)
        missing = [p for p in slides(pid)["items"] if p.get("keep") and not p.get("image_url")]
        if missing:
            raise HTTPException(400, "仍有 %d 页结果未就绪；暂不支持部分图片版导出" % len(missing))
    if kind in ("image", "editable_hi"):
        pending = [p for p in slides(pid)["items"] if p.get("keep") and p["review_required"]]
        if pending:
            raise HTTPException(400, "生成图片待核对，请逐页对照原文并确认后导出")
    if kind in ('beautify_editable', 'beautify_image'):
        if not db.q1('SELECT id FROM projects WHERE id=?', (pid,)):
            raise HTTPException(404, '项目不存在')
        state = beautify.read_state(pid)
        proj = _beautify_project(pid)
        if state.get('source') != [beautify.fingerprint(proj['source_path'])]:
            raise HTTPException(400, '原稿已变更，请重新生成美化预览')
        key = 'editable' if kind == 'beautify_editable' else 'image'
        path = state.get('exports', {}).get(key, '')
        if not path or not os.path.isfile(path):
            raise HTTPException(400, '暂无美化产物，请先跑美化；可编辑版需精细重建成功')
        return FileResponse(path, filename=re.sub(r"[^\w.\u4e00-\u9fff-]", "_",
                            (proj.get("title") or "deck") + "_" + kind + ".pptx"),
                            media_type='application/vnd.openxmlformats-officedocument.presentationml.presentation')
    if kind == "editable":
        raise HTTPException(400, "「本地渲染的可编辑导出」随内置引擎一并下线："
                                 "请走「图片转可编辑」模块（#/editable）用第三方 skill 精细重建，"
                                 "产物用 kind=editable_hi 下载")
    if kind not in ("image", "editable_hi"):
        raise HTTPException(400, "kind 只能是 image 或 editable_hi")
    if kind == "editable_hi":
        row = db.q1("SELECT * FROM exports WHERE project_id=? AND kind='editable_hi' "
                    "ORDER BY created_at DESC LIMIT 1", (pid,))
        if not row:
            raise HTTPException(400, "还没有转换产物；请先在「图片转可编辑」里转换页面")
        data = dict(row)
        if not os.path.exists(data.get("path") or ""):
            raise HTTPException(400, "产物文件已不在磁盘上，请重新转换")
        project = get_project(pid)
        name = re.sub(r"[^\w.\u4e00-\u9fff-]", "_", project["title"] + "_editable_hi_" + data["id"] + ".pptx")
        renamed = os.path.join(os.path.dirname(data["path"]), name)
        if os.path.abspath(data["path"]) != os.path.abspath(renamed):
            import shutil
            shutil.copyfile(data["path"], renamed)
            data["path"] = renamed
        pages = json.loads(data.get("editable_pages_json") or "[]")
        return dict(kind="editable_hi", engine="editppt",
                    file="/files/" + os.path.relpath(data["path"], config.FILES_DIR).replace(os.sep, "/"),
                    path=data["path"], pages=data.get("pages") or len(pages), editable_pages=pages,
                    size_bytes=data.get("size_bytes") or 0, created_at=data.get("created_at") or "",
                    note="第三方 image-to-editable-ppt skill 的精细重建产物，**仅含已转换页面**")
    try:
        return export_snapshots.create(pid, export_preview(pid, {"kind": "image"}))
    except ValueError as e:
        raise HTTPException(400, str(e))


# ────────────────────────── 作业（SSE）──────────────────────────

@app.get("/api/jobs/{job_id}")
def job_detail(job_id: str) -> Dict[str, Any]:
    job = pipeline.get_job(job_id)
    if not job:
        raise HTTPException(404, "作业不存在")
    return job


@app.post("/api/jobs/{job_id}/cancel")
async def cancel_job(job_id: str) -> Dict[str, Any]:
    job = pipeline.get_job(job_id)
    if not job:
        raise HTTPException(404, "作业不存在")
    if job['status'] in ('done', 'failed', 'cancelled', 'interrupted'):
        raise HTTPException(400, "作业已结束，当前状态：%s" % job['status'])
    pipeline.request_cancel(job_id)
    pipeline._job_update(job_id, status='cancelling', message='正在取消')
    return dict(ok=True, job_id=job_id, status='cancelling', message='正在取消')


def _job_terminal_event(job):
    status = job.get('status')
    return dict(type='interrupted' if status == 'interrupted' else 'cancelled' if status == 'cancelled' else 'error' if status == 'failed' else 'done',
                job_id=job.get('id'), version=job.get('version', 0), message='已取消生成' if status == 'cancelled' else job.get('message', ''), job=job)


@app.get("/api/jobs/{job_id}/stream")
async def job_stream(job_id: str) -> StreamingResponse:
    job = pipeline.get_job(job_id)
    if not job:
        raise HTTPException(404, "作业不存在")
    async def gen():
        q = pipeline.stream_open(job_id)
        try:
            if q is None:
                if job.get("status") in ("done", "failed", "cancelled", "interrupted"):
                    yield "data: %s\n\n" % json.dumps(_job_terminal_event(job), ensure_ascii=False)
                else:
                    # A persisted active task without a live worker is unknown, never success.
                    yield ": 执行状态待恢复，请查询同一任务\n\n"
                return
            while True:
                try:
                    ev = await asyncio.wait_for(q.get(), timeout=20.0)
                except asyncio.TimeoutError:
                    yield ": ping\n\n"                     # 心跳，防代理断流
                    cur = pipeline.get_job(job_id) or {}
                    if cur.get("status") in ("done", "failed", "cancelled", "interrupted"):
                        yield "data: %s\n\n" % json.dumps(_job_terminal_event(cur), ensure_ascii=False)
                        return
                    continue
                yield "data: %s\n\n" % json.dumps(ev, ensure_ascii=False)
                if ev.get("type") in ("done", "error", "cancelled", "interrupted"):
                    return
        finally:
            pipeline.stream_close(job_id)

    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


# ────────────────────────── 用量 ──────────────────────────

@app.get("/api/provider-usage")
def provider_usage_endpoint(refresh: bool = Query(False)) -> Dict[str, Any]:
    """供应商面板侧用量（New API 兼容计费接口，当前 GuaiHub）。

    只读、不花钱；60 秒缓存。失败时返回 `ok:false + error`，不抛 5xx —— 统计页绝不因此打不开。
    """
    return provider_usage.summary(refresh=refresh)


@app.get("/api/usage")
def usage(range: str = Query("month")) -> Dict[str, Any]:
    where = "" if range == "all" else " WHERE ts >= ?"
    args = () if range == "all" else (
        datetime.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0).isoformat(),)
    rows = [dict(r) for r in db.q(
        "SELECT * FROM usage_log" + where + " ORDER BY datetime(ts) DESC LIMIT 500", args)]
    total = round(float(db.q1(
        "SELECT COALESCE(SUM(cny), 0) AS total FROM usage_log" + where, args)["total"]), 4)

    def agg(field: str) -> List[Dict[str, Any]]:
        return [dict(**{field: r[field]}, cny=round(float(r["cny"]), 4), count=r["count"])
                for r in db.q(
                    "SELECT COALESCE(NULLIF(" + field + ", ''), '-') AS " + field +
                    ", COALESCE(SUM(cny), 0) AS cny, COUNT(*) AS count FROM usage_log" +
                    where + " GROUP BY COALESCE(NULLIF(" + field + ", ''), '-') ORDER BY cny DESC", args)]

    by_project = [dict(r) for r in db.q(
        "SELECT project_id, COALESCE(SUM(cny),0) AS cny, COUNT(*) AS count FROM usage_log" +
        where + " GROUP BY project_id ORDER BY cny DESC", args)]
    for item in by_project:
        project = db.q1("SELECT title FROM projects WHERE id=?", (item["project_id"],))
        item["title"] = project["title"] if project else "已删除项目" if item["project_id"] else "未关联项目"
    return dict(total_cny=total, by_action=agg("action"), by_provider=agg("provider"), by_project=by_project, items=rows)


# ────────────────────────── PPT 备注（上传已有 pptx → 生成讲稿备注 → 写回备注栏）──────────────────────────

@app.get("/api/notes")
def notes_list() -> Dict[str, Any]:
    rows = [dict(r) for r in db.q("SELECT * FROM notes_jobs ORDER BY datetime(created_at) DESC LIMIT 200")]
    for r in rows:
        r["cost_cny"] = float(r.get("cost_cny") or 0.0)
    return dict(items=rows)


@app.post("/api/notes/upload")
async def notes_upload(file: UploadFile = File(...)) -> Dict[str, Any]:
    name = file.filename or "deck.pptx"
    ext = os.path.splitext(name)[1].lower()
    if ext != ".pptx":
        raise HTTPException(400, "只支持 .pptx 文件。%s 旧格式请先在 PowerPoint 里另存为 .pptx"
                            % (ext or "该"))
    nid = db.uid("n")
    dest_dir = os.path.join(config.FILES_DIR, "notes", nid)
    os.makedirs(dest_dir, exist_ok=True)
    dest = os.path.join(dest_dir, "source.pptx")
    raw = await file.read()
    if not raw:
        raise HTTPException(400, "文件内容为空")
    with open(dest, "wb") as f:
        f.write(raw)
    try:
        detail = await asyncio.to_thread(slide_notes.create_job, dest, name)
    except Exception as e:                               # noqa: BLE001
        raise HTTPException(400, "解析 PPT 失败：%s" % str(e)[:300])
    if not detail.get("pages"):
        raise HTTPException(400, "这个 PPT 没有可读取的幻灯片")
    detail["size_bytes"] = len(raw)
    detail["hint"] = ("其中 %d 页文字很少（图片页），需要画面识别才能写备注" % detail.get("image_pages", 0)) \
        if detail.get("image_pages") else "页面文字充足，可只用文本框内容生成"
    return detail


@app.get("/api/notes/{nid}")
def notes_detail(nid: str) -> Dict[str, Any]:
    detail = slide_notes.job_detail(nid)
    if not detail:
        raise HTTPException(404, "备注任务不存在")
    return detail


@app.post("/api/notes/{nid}/render")
def notes_render(nid: str) -> Dict[str, Any]:
    """按需把上传的 pptx 渲染成逐页 PNG（幂等）：逐页复核里显示 PPT 实际页面用。

    首次调用跑一次 LibreOffice（约 10~30 秒），之后走磁盘缓存；页数不一致抛 400（防错页）。
    """
    if not slide_notes.job_row(nid):
        raise HTTPException(404, "备注任务不存在")
    try:
        return slide_notes.ensure_rendered(nid)
    except ValueError as e:
        raise HTTPException(400, str(e))


@app.delete("/api/notes/{nid}")
def notes_delete(nid: str) -> Dict[str, Any]:
    if not slide_notes.job_row(nid):
        raise HTTPException(404, "备注任务不存在")
    db.x("DELETE FROM notes_pages WHERE job_id=?", (nid,))
    db.x("DELETE FROM notes_jobs WHERE id=?", (nid,))
    return dict(ok=True)


@app.post("/api/notes/{nid}/analyze")
async def notes_analyze(nid: str, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    if not slide_notes.job_row(nid):
        raise HTTPException(404, "备注任务不存在")
    mode = payload.get("vision_mode") or "auto"
    try:
        job_id = await slide_notes.run_analyze(nid, mode, payload.get("page_nos"))
    except ValueError as e:
        raise HTTPException(400, str(e))
    return dict(job_id=job_id, vision_mode=mode, progress_url="/api/jobs/%s/stream" % job_id,
                estimate=dict(pages=len(payload.get("page_nos") or []) or None,
                              note="画面识别每页约 5~110 秒（页面越复杂越慢），3 页并发"))


@app.post("/api/notes/{nid}/generate")
async def notes_generate(nid: str, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    if not slide_notes.job_row(nid):
        raise HTTPException(404, "备注任务不存在")
    try:
        job_id = await slide_notes.run_generate(nid, payload.get("page_nos"),
                                                str(payload.get("instruction") or ""),
                                                bool(payload.get("only_missing")))
    except ValueError as e:
        raise HTTPException(400, str(e))
    return dict(job_id=job_id, progress_url="/api/jobs/%s/stream" % job_id)


@app.post("/api/notes/{nid}/pages/{page_no}/regenerate")
async def notes_regenerate(nid: str, page_no: int, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    if not slide_notes.job_row(nid):
        raise HTTPException(404, "备注任务不存在")
    try:
        job_id = await slide_notes.run_generate(nid, [page_no], str(payload.get("instruction") or ""))
    except ValueError as e:
        raise HTTPException(400, str(e))
    return dict(job_id=job_id, progress_url="/api/jobs/%s/stream" % job_id)


@app.patch("/api/notes/{nid}/pages/{page_no}")
def notes_save_page(nid: str, page_no: int, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    row = db.q1("SELECT * FROM notes_pages WHERE job_id=? AND page_no=?", (nid, page_no))
    if not row:
        raise HTTPException(404, "该页不存在")
    note = payload.get("note")
    if not isinstance(note, str):
        raise HTTPException(400, "note 必须是字符串")
    upd = dict(note=note.strip(), note_edited=1, note_status="done" if note.strip() else "pending")
    if upd["note_status"] == "done":
        upd["error"] = ""
    db.update("notes_pages", row["id"], upd)
    return dict(ok=True, page_no=page_no, note=upd["note"], note_status=upd["note_status"])


@app.post("/api/notes/{nid}/export")
def notes_export(nid: str, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    if not slide_notes.job_row(nid):
        raise HTTPException(404, "备注任务不存在")
    keep = payload.get("keep_original")
    try:
        return slide_notes.export_with_notes(nid, keep_original=True if keep is None else bool(keep))
    except ValueError as e:
        raise HTTPException(400, str(e))


# ────────────────────────── 风格构建（上传模板/参考图 → 自动解析 → 落成新风格包）──────────────────────────

@app.post("/api/style-builder/upload")
async def style_builder_upload(files: Optional[List[UploadFile]] = File(None),
                               file: Optional[UploadFile] = File(None)) -> Dict[str, Any]:
    """上传模板（.pptx）或风格参考图（可多张）。返回 builder_id，随后可 analyze。"""
    incoming = [f for f in (files or []) if f is not None]
    if file is not None:
        incoming.append(file)
    if not incoming:
        raise HTTPException(400, "没有收到文件（支持 .pptx / .pdf / .png / .jpg / .jpeg / .webp）")
    pairs: List[Any] = []
    for item in incoming:
        raw = await item.read()
        if raw:
            pairs.append((item.filename or "file", raw))
    try:
        out = style_builder.save_uploads(pairs)
    except ValueError as e:
        raise HTTPException(400, str(e))
    # PPTX：立刻后台转代表页（与 analyze 阶段共用同一份缓存，解析时直接命中）
    if out.get("kind") in ("pptx", "pdf") and style_builder.start_pages(out["builder_id"]):
        asyncio.create_task(style_builder.build_pages_background(out["builder_id"]))
    return out


@app.post("/api/style-builder/{bid}/analyze")
async def style_builder_analyze(bid: str, payload: Optional[Dict[str, Any]] = Body(None)) -> Dict[str, Any]:
    """开始解析（后台作业，进度走既有 /api/jobs/{id}/stream）。"""
    body = payload or {}
    try:
        job_id = await style_builder.run_analyze(bid, bool(body.get("use_vision", True)), body.get("pages"))
    except ValueError as e:
        raise HTTPException(404, str(e))
    return dict(job_id=job_id, progress_url="/api/jobs/%s/stream" % job_id)


@app.get("/api/style-builder/{bid}")
def style_builder_get(bid: str) -> Dict[str, Any]:
    out = style_builder.response(bid)
    if not out:
        raise HTTPException(404, "上传记录不存在")
    return out


@app.post("/api/style-builder/{bid}/save")
def style_builder_save(bid: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    """把预览结果落成 style-packs/styles/<id>/（同名 id 一律拒绝，不覆盖既有风格包）。"""
    try:
        return style_builder.save(bid, payload)
    except ValueError as e:
        raise HTTPException(400, str(e))


@app.delete("/api/style-builder/{bid}")
def style_builder_delete(bid: str) -> Dict[str, Any]:
    try:
        return style_builder.delete(bid)
    except ValueError as e:
        raise HTTPException(404, str(e))


def _beautify_project(pid):
    proj = db.q1('SELECT * FROM projects WHERE id=?', (pid,))
    if not proj:
        raise HTTPException(404, '项目不存在')
    src = proj['source_path'] or ''
    if not src.lower().endswith(('.pptx', '.pdf')) or not os.path.isfile(src):
        raise HTTPException(400, '请上传 PPTX 或 PDF 原稿')
    return dict(proj)


@app.post('/api/projects/{pid}/beautify/origin')
async def beautify_origin(pid: str) -> Dict[str, Any]:
    try:
        proj = _beautify_project(pid)
        rendered = await asyncio.to_thread(beautify.render_origin, pid, proj['source_path'])
    except HTTPException as exc:
        raise HTTPException(400, exc.detail)
    return dict(page_count=rendered['page_count'], renderer=rendered['renderer'],
                pages=[dict(page_no=p['page_no'], image_url=pipeline._files_url(p['path']),
                            width=p['width'], height=p['height']) for p in rendered['pages']])


async def _beautify_selection(proj, payload):
    count = await asyncio.to_thread(beautify.page_count, proj['source_path'])
    try:
        return beautify.validate_page_nos(payload.get('page_nos'), count, 'page_nos' in payload)
    except ValueError as exc:
        raise HTTPException(400, str(exc))


def _guard_beautify_content(pid):
    if db.q1("SELECT id FROM jobs WHERE project_id=? AND kind='beautify' AND status IN ('queued','running','cancelling')", (pid,)):
        raise HTTPException(409, '美化作业运行中，暂不能修改内容')


@app.post('/api/projects/{pid}/beautify/content')
async def beautify_content(pid: str, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    proj = _beautify_project(pid)
    _guard_beautify_content(pid)
    try:
        result = await asyncio.to_thread(beautify.prepare_content, pid, proj['source_path'],
                                        payload.get('page_nos'), payload.get('keep_material_page_nos', []))
    except (ValueError, TypeError, IndexError) as exc:
        raise HTTPException(400, str(exc))
    _guard_beautify_content(pid)
    return result


@app.put('/api/projects/{pid}/beautify/content')
def beautify_content_save(pid: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    proj = _beautify_project(pid)
    with beautify.CONTENT_LOCK:
        _guard_beautify_content(pid)
        try:
            return beautify.save_content(pid, proj['source_path'], payload)
        except (ValueError, TypeError, IndexError) as exc:
            raise HTTPException(400, str(exc))


@app.post('/api/projects/{pid}/beautify/preview')
async def beautify_preview(pid: str, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    proj = _beautify_project(pid)
    page_nos = await _beautify_selection(proj, payload)
    mode = str(payload.get('mode', 'img2img'))
    if mode not in ('img2img', 'redesign'):
        raise HTTPException(400, 'mode 只能是 img2img 或 redesign')
    try:
        job = await pipeline.run_beautify(pid, str(payload.get('style_pack_id') or ''), 'preview',
                     mode, int(payload.get('preview_pages', 3)), str(payload.get('size', '2K')), page_nos=page_nos)
    except (ValueError, TypeError) as exc:
        raise HTTPException(400, str(exc))
    return dict(job_id=job, progress_url=f'/api/jobs/{job}/stream',
                page_count=await asyncio.to_thread(beautify.page_count, proj['source_path']), message='已开始预览美化')


@app.post('/api/projects/{pid}/beautify/pages/{page_no}/regenerate')
async def beautify_regenerate_page(pid: str, page_no: int, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    """只重生成指定一页：非破坏性，其它页的美图与导出不动。"""
    _beautify_project(pid)
    try:
        job = await pipeline.run_beautify_page(pid, page_no, str(payload.get('style_pack_id') or ''),
                                               str(payload.get('mode') or ''), str(payload.get('size') or '2K'),
                                               str(payload.get('instruction') or ''))
    except (ValueError, TypeError) as exc:
        raise HTTPException(400, str(exc))
    return dict(job_id=job, progress_url=f'/api/jobs/{job}/stream', page_no=page_no,
                message='已开始按你的修改要求重新设计该页' if (payload.get('instruction') or '').strip() else '已开始重生成该页')


@app.get('/api/projects/{pid}/beautify/pages/{page_no}/versions')
def beautify_versions(pid: str, page_no: int) -> Dict[str, Any]:
    """该页历史版本列表：当前正在使用的图单独给 `current`，不在 versions 里（前端标注「当前使用」）。"""
    _beautify_project(pid)
    return beautify.list_versions(pid, page_no)


@app.post('/api/projects/{pid}/beautify/pages/{page_no}/versions/{version_id}/restore')
async def beautify_restore(pid: str, page_no: int, version_id: str) -> Dict[str, Any]:
    """把该页恢复到某个历史版本（恢复前会先归档当前图，可再切回）。"""
    _beautify_project(pid)
    try:
        job = await pipeline.run_beautify_restore(pid, page_no, version_id)
    except (ValueError, TypeError) as exc:
        raise HTTPException(400, str(exc))
    return dict(job_id=job, progress_url=f'/api/jobs/{job}/stream', page_no=page_no, message='已开始恢复该页历史版本')


@app.post('/api/projects/{pid}/beautify/continue')
async def beautify_continue(pid: str, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    proj = _beautify_project(pid)
    page_nos = await _beautify_selection(proj, payload)
    state = beautify.read_state(pid)
    if state.get('phase') != 'preview':
        raise HTTPException(400, '请先完成预览')
    mode = str(payload.get('mode', state.get('mode', 'img2img')))
    if mode not in ('img2img', 'redesign'):
        raise HTTPException(400, 'mode 只能是 img2img 或 redesign')
    remaining = len(set(page_nos) - {p['page_no'] for p in state.get('pages', []) if os.path.isfile(p.get('beauty') or '')})
    try:
        job = await pipeline.run_beautify(pid, str(payload.get('style_pack_id') or state.get('style_pack_id', '')),
                     'full', mode, size=str(payload.get('size', '2K')), page_nos=page_nos)
    except (ValueError, TypeError) as exc:
        raise HTTPException(400, str(exc))
    return dict(job_id=job, progress_url=f'/api/jobs/{job}/stream', remaining_pages=remaining)


@app.post('/api/projects/{pid}/beautify/editable')
async def beautify_editable(pid: str, payload: Dict[str, Any] = Body(default={})) -> Dict[str, Any]:
    """按需把已生成的美化图片版转成原生可编辑版（生成图片时不再自动执行）。"""
    _beautify_project(pid)
    try:
        job = await pipeline.run_beautify_editable(pid, str(payload.get('style_pack_id') or ''))
    except (ValueError, TypeError) as exc:
        raise HTTPException(400, str(exc))
    return dict(job_id=job, progress_url=f'/api/jobs/{job}/stream', message='已开始重建可编辑版')


@app.get('/api/projects/{pid}/beautify')
def beautify_status(pid: str) -> Dict[str, Any]:
    proj = _beautify_project(pid)
    state = beautify.read_state(pid)
    if state and state.get('source') != [beautify.fingerprint(proj['source_path'])]:
        state = dict(warnings=['原稿已变更，请重新生成美化预览'])
    result = dict(phase=state.get('phase', ''), mode=state.get('mode', 'img2img'), page_count=state.get('page_count', 0),
                  page_nos=state.get('page_nos', []), rendered=state.get('rendered', 0), confirmed=state.get('confirmed', False),
                  pages=[], editable_engine=state.get('editable_engine'), exports={}, warnings=state.get('warnings', []),
                  concurrency=pipeline.BEAUTIFY_CONCURRENCY)
    for page in state.get('pages', []):
        result['pages'].append(dict(page_no=page['page_no'],
            origin_url=pipeline._files_url(page['origin']) if os.path.isfile(page.get('origin', '')) else '',
            beauty_url=pipeline._files_url(page['beauty']) if os.path.isfile(page.get('beauty', '')) else '',
            beauty_ok=page.get('beauty_ok', False), text_source=page.get('text_source', 'none'),
            keep_material=bool(page.get('keep_material', False)),
            preserved_pictures=int(page.get('preserved_pictures') or 0),
            instruction=str(page.get('instruction') or ''),
            version_count=int(page.get('versions') or 0),
            restored_from=str(page.get('restored_from') or '')))
    result['exports'] = {key: pipeline._files_url(path) for key, path in state.get('exports', {}).items() if os.path.isfile(path)}
    # 内容是否**已保存并确认**（磁盘上存在与当前原稿匹配的 content.json）——恢复流程时用它判断第③步是否已完成
    result['content_saved'] = bool(beautify.content_plan(pid, proj['source_path']))
    result['updated_at'] = state.get('updated_at') or ''
    result['cost'] = beautify.cost_summary(pid)
    return result


@app.post("/api/projects/{pid}/quote")
def operation_quote(pid: str, payload: Dict[str, Any] = Body(...)) -> Dict[str, Any]:
    return operations.quote(pid, payload)


@app.get("/api/projects/{pid}/submissions/{key}")
def operation_lookup(pid: str, key: str) -> Dict[str, Any]:
    return operations.lookup(pid, key)


@app.get("/api/projects/{pid}/usage")
def project_usage(pid: str) -> Dict[str, Any]:
    return operations.ledger(pid)


# ⚠️ 静态挂载必须位于所有 API 端点之后：app.mount("/") 会吞掉其后注册的路由（踩过坑：美化 v2 的端点曾因此全部 404）。
# ────────────────────────── 静态文件与前端 ──────────────────────────

app.mount("/files", StaticFiles(directory=config.FILES_DIR), name="files")
# 风格包的封面/预览图（style-packs/styles/<id>/previews/*.png）
app.mount(stylepacks.ASSET_URL, StaticFiles(directory=os.path.join(config.STYLE_DIR, "styles")),
          name="style-assets")

WEB_DIST = os.path.join(config.ROOT, "web", "dist")


class WebStaticFiles(StaticFiles):
    """前端静态资源：HTML 不缓存，其余交给默认行为。

    前端包名带内容哈希（index-<hash>.js），新构建自然会换名字；真正会让人
    「明明改了却还看到旧界面」的是 index.html 被浏览器按启发式缓存住 ——
    所以只给 HTML 加 no-cache（仍带 ETag，命中时走 304，开销可忽略）。
    """

    async def get_response(self, path, scope):
        response = await super().get_response(path, scope)
        ctype = response.headers.get("content-type", "")
        if ctype.startswith("text/html"):
            response.headers["Cache-Control"] = "no-cache"
        return response


if os.path.isdir(WEB_DIST):
    app.mount("/", WebStaticFiles(directory=WEB_DIST, html=True), name="web")
else:
    @app.get("/")
    def index() -> JSONResponse:
        return JSONResponse(dict(
            name="SlideCraft API", version=VERSION,
            frontend="未构建：请在 web/ 执行 npm run build（或 npm run dev 走 :5183）",
            docs="/docs", api_contract="docs/API.md"))
